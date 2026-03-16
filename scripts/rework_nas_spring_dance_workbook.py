from __future__ import annotations

import argparse
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


SESSION_12 = "12:00–1:30 PM"
SESSION_2 = "2–3:30 PM"

OUTPUT_HEADERS = [
    "Student Last Name",
    "Student First Name",
    "Final Session",
    "Grade",
    "Teacher",
    "Payer Name",
    "Payer Email",
    "Payment Status",
    "Adjustment Note",
]

REQUEST_LOG_HEADERS = [
    "Student Last Name",
    "Student First Name",
    "Payer Name",
    "Payer Email",
    "Original Session",
    "Final Session",
    "Payment Status",
    "Request Type",
    "Workbook Action",
    "Request Note",
]


@dataclass(frozen=True)
class SessionOverride:
    student_first: str
    student_last: str
    payer_email: str
    from_session: str
    to_session: str
    note: str


@dataclass(frozen=True)
class RequestEntry:
    student_first: str
    student_last: str
    payer_email: str
    payer_name: str
    request_type: str
    workbook_action: str
    request_note: str
    match_payer_email: str | None = None


OVERRIDES = {
    ("christopher", "saulino", "marykate.saulino@gmail.com"): SessionOverride(
        student_first="Christopher",
        student_last="Saulino",
        payer_email="marykate.saulino@gmail.com",
        from_session=SESSION_2,
        to_session=SESSION_12,
        note="Moved to 12:00–1:30 PM per Mary Kate Saulino switch-time email screenshot.",
    ),
    ("john", "jenkins", "jennajenkins0610@yahoo.com"): SessionOverride(
        student_first="John",
        student_last="Jenkins",
        payer_email="jennajenkins0610@yahoo.com",
        from_session=SESSION_2,
        to_session=SESSION_12,
        note="Moved to 12:00–1:30 PM per Jay/Jenkins switch-time email screenshot.",
    ),
    ("chelsea", "cathel", "pcathel+newalbany@gmail.com"): SessionOverride(
        student_first="Chelsea",
        student_last="Cathel",
        payer_email="pcathel+newalbany@gmail.com",
        from_session=SESSION_2,
        to_session=SESSION_12,
        note="Moved to 12:00–1:30 PM per Hannah Cathel switch-time email screenshot.",
    ),
    ("andrey", "myketey", "lindsaylee.morgan@gmail.com"): SessionOverride(
        student_first="Andrey",
        student_last="Myketey",
        payer_email="lindsaylee.morgan@gmail.com",
        from_session=SESSION_2,
        to_session=SESSION_12,
        note="Moved to 12:00–1:30 PM per Lindsay Lee Myketey message screenshot.",
    ),
    ("remy", "gleason", "sdsugleason@gmail.com"): SessionOverride(
        student_first="Remy",
        student_last="Gleason",
        payer_email="sdsugleason@gmail.com",
        from_session=SESSION_2,
        to_session=SESSION_12,
        note="Moved to 12:00–1:30 PM per Megan Gleason switch-time email screenshot.",
    ),
}

REQUEST_LOG = [
    RequestEntry(
        student_first="Christopher",
        student_last="Saulino",
        payer_email="marykate.saulino@gmail.com",
        payer_name="Mary Kate Saulino",
        match_payer_email=None,
        request_type="Session swap request",
        workbook_action="Moved from 2:00 PM session to 12:00 PM session.",
        request_note="Mary Kate Saulino asked to move Christopher Saulino from 2:00 PM to 12:00 PM.",
    ),
    RequestEntry(
        student_first="John",
        student_last="Jenkins",
        payer_email="jennajenkins0610@yahoo.com",
        payer_name="Jenna Jenkins",
        match_payer_email=None,
        request_type="Session swap request",
        workbook_action="Moved from 2:00 PM session to 12:00 PM session.",
        request_note="Jay/Jenkins request was matched to John Jenkins and moved from 2:00 PM to 12:00 PM.",
    ),
    RequestEntry(
        student_first="Chelsea",
        student_last="Cathel",
        payer_email="pcathel+newalbany@gmail.com",
        payer_name="Paul-Michael Cathel",
        match_payer_email=None,
        request_type="Session swap request",
        workbook_action="Moved from 2:00 PM session to 12:00 PM session.",
        request_note="Hannah Cathel asked to move Chelsea Cathel from 2:00 PM to 12:00 PM.",
    ),
    RequestEntry(
        student_first="Bodhi",
        student_last="Lepore",
        payer_email="hetalpateleco@gmail.com",
        payer_name="Hetal Patel",
        match_payer_email="",
        request_type="Payment / registration follow-up",
        workbook_action="No attendee list change needed.",
        request_note="Hetal Patel asked whether Bodhi Lepore's form and payment were received. Bodhi was already on the 12:00 PM attendee list.",
    ),
    RequestEntry(
        student_first="Andrey",
        student_last="Myketey",
        payer_email="lindsaylee.morgan@gmail.com",
        payer_name="Lindsay Lee Myketey",
        match_payer_email=None,
        request_type="Session swap request",
        workbook_action="Moved from 2:00 PM session to 12:00 PM session.",
        request_note="Lindsay Lee Myketey asked to move Dre/Andrey Myketey to the earlier session.",
    ),
    RequestEntry(
        student_first="Remy",
        student_last="Gleason",
        payer_email="sdsugleason@gmail.com",
        payer_name="Megan Gleason",
        match_payer_email=None,
        request_type="Session swap request",
        workbook_action="Moved from 2:00 PM session to 12:00 PM session.",
        request_note="Megan Gleason asked to move Remy Gleason to the earlier session.",
    ),
    RequestEntry(
        student_first="Henry",
        student_last="Rizzolo",
        payer_email="chelsea.caro@gmail.com",
        payer_name="Chelsea Caro",
        match_payer_email=None,
        request_type="Late registration request",
        workbook_action="Kept on 2:00 PM attendee list with pending payment status.",
        request_note="Chelsea Caro asked whether Henry Rizzolo could still register for 2:00 PM. The workbook already had Henry listed at 2:00 PM with pending payment status.",
    ),
]

ADJUSTMENT_NOTES = {
    ("bodhi", "lepore", ""): (
        "Parent follow-up received; Bodhi Lepore was already listed at 12:00–1:30 PM, so no attendee change was needed."
    ),
    ("henry", "rizzolo", "chelsea.caro@gmail.com"): (
        "Registration request received; Henry Rizzolo remains on the 2–3:30 PM list with pending payment status."
    ),
}


def normalize_text(value: object) -> str:
    if value is None:
        return ""
    return " ".join(str(value).split())


def normalize_key(*parts: object) -> tuple[str, ...]:
    return tuple(normalize_text(part).casefold() for part in parts)


def sheet_row_to_dict(sheet) -> Iterable[dict[str, object]]:
    rows = sheet.iter_rows(values_only=True)
    headers = [normalize_text(cell) for cell in next(rows)]
    for row in rows:
        yield {headers[index]: row[index] for index in range(len(headers))}


def build_output_row(record: dict[str, object]) -> list[str]:
    payer_name = " ".join(
        part for part in [record["Payer First Name"], record["Payer Last Name"]] if part
    )
    return [
        normalize_text(record["New Albany Student Last Name"]),
        normalize_text(record["New Albany Student First Name"]),
        normalize_text(record["Final Session"]),
        normalize_text(record["Student's Grade"]),
        normalize_text(record["Teacher's Name"]),
        normalize_text(payer_name),
        normalize_text(record["Payer Email"]),
        normalize_text(record["Payment Status"]),
        normalize_text(record["Adjustment Note"]),
    ]


def build_request_log_row(
    request: RequestEntry,
    original_session: str,
    final_session: str,
    payment_status: str,
) -> list[str]:
    return [
        request.student_last,
        request.student_first,
        request.payer_name,
        request.payer_email,
        original_session,
        final_session,
        payment_status,
        request.request_type,
        request.workbook_action,
        request.request_note,
    ]


def sort_key(row: list[str]) -> tuple[str, str, str, str]:
    return (
        row[0].casefold(),
        row[1].casefold(),
        row[2].casefold(),
        row[5].casefold(),
    )


def remove_sheet_if_present(workbook, title: str) -> None:
    if title in workbook.sheetnames:
        del workbook[title]


def format_sheet(sheet) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    header_font = Font(bold=True)
    for cell in sheet[1]:
        cell.font = header_font

    for column_index, column_cells in enumerate(sheet.iter_cols(), start=1):
        width = 0
        for cell in column_cells:
            width = max(width, len(normalize_text(cell.value)))
        sheet.column_dimensions[get_column_letter(column_index)].width = min(width + 2, 48)


def create_output_sheet(workbook, title: str, rows: list[list[str]]) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(OUTPUT_HEADERS)
    for row in rows:
        sheet.append(row)
    format_sheet(sheet)


def derive_destination(source_path: Path) -> Path:
    return source_path.with_name(f"{source_path.stem}_reworked.xlsx")


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Rework the NAS Spring Dance 2026 CheddarUp export into attendee tabs."
    )
    parser.add_argument("source", type=Path, help="Path to the original collection summary workbook.")
    parser.add_argument(
        "--dest",
        type=Path,
        help="Optional destination path for the updated workbook. Defaults to '<source>_reworked.xlsx'.",
    )
    args = parser.parse_args()

    source_path = args.source.resolve()
    if not source_path.exists():
        raise FileNotFoundError(f"Workbook not found: {source_path}")

    destination_path = (args.dest.resolve() if args.dest else derive_destination(source_path))

    workbook = load_workbook(source_path)
    source_sheet = workbook["1-EntryTicket"]

    canonical_rows: list[list[str]] = []
    applied_overrides: list[str] = []
    request_index = {
        normalize_key(
            entry.student_first,
            entry.student_last,
            entry.payer_email if entry.match_payer_email is None else entry.match_payer_email,
        ): entry
        for entry in REQUEST_LOG
    }
    request_log_rows: list[list[str]] = []
    captured_request_keys: set[tuple[str, ...]] = set()

    for record in sheet_row_to_dict(source_sheet):
        payment_status = normalize_text(record["Payment Status"])
        if payment_status.casefold() == "refunded":
            continue

        student_first = normalize_text(record["New Albany Student First Name"])
        student_last = normalize_text(record["New Albany Student Last Name"])
        payer_email = normalize_text(record["Payer Email"])
        current_session = normalize_text(record["Time"])
        override_key = normalize_key(student_first, student_last, payer_email)
        override = OVERRIDES.get(override_key)

        final_session = current_session
        adjustment_note = ""
        if override:
            if current_session != override.from_session:
                raise ValueError(
                    f"Expected {student_first} {student_last} to be in '{override.from_session}', "
                    f"but found '{current_session}'."
                )
            final_session = override.to_session
            adjustment_note = override.note
            applied_overrides.append(f"{student_first} {student_last}")
        elif override_key in ADJUSTMENT_NOTES:
            adjustment_note = ADJUSTMENT_NOTES[override_key]

        enriched_record = dict(record)
        enriched_record["Final Session"] = final_session
        enriched_record["Adjustment Note"] = adjustment_note
        canonical_rows.append(build_output_row(enriched_record))

        request_entry = request_index.get(override_key)
        if request_entry:
            request_log_rows.append(
                build_request_log_row(
                    request=request_entry,
                    original_session=current_session,
                    final_session=final_session,
                    payment_status=payment_status,
                )
            )
            captured_request_keys.add(override_key)

    canonical_rows.sort(key=sort_key)
    noon_rows = [row for row in canonical_rows if row[2] == SESSION_12]
    two_pm_rows = [row for row in canonical_rows if row[2] == SESSION_2]

    missing_requests = [
        entry
        for entry in REQUEST_LOG
        if normalize_key(
            entry.student_first,
            entry.student_last,
            entry.payer_email if entry.match_payer_email is None else entry.match_payer_email,
        )
        not in captured_request_keys
    ]
    if missing_requests:
        missing_names = ", ".join(f"{entry.student_first} {entry.student_last}" for entry in missing_requests)
        raise ValueError(f"Could not match request log entries in workbook: {missing_names}")

    request_log_rows.sort(key=sort_key)

    for title in ("Master by Student", "12PM Attendees", "2PM Attendees", "Request Log"):
        remove_sheet_if_present(workbook, title)

    create_output_sheet(workbook, "Master by Student", canonical_rows)
    create_output_sheet(workbook, "12PM Attendees", noon_rows)
    create_output_sheet(workbook, "2PM Attendees", two_pm_rows)
    request_sheet = workbook.create_sheet("Request Log")
    request_sheet.append(REQUEST_LOG_HEADERS)
    for row in request_log_rows:
        request_sheet.append(row)
    format_sheet(request_sheet)

    workbook.save(destination_path)

    print(f"Created workbook: {destination_path}")
    print(f"Included attendees: {len(canonical_rows)}")
    print(f"{SESSION_12}: {len(noon_rows)}")
    print(f"{SESSION_2}: {len(two_pm_rows)}")
    print(f"Tracked requests: {len(request_log_rows)}")
    print("Applied overrides:")
    for name in sorted(applied_overrides):
        print(f"- {name}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
