from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches as DocxInches, Pt, RGBColor as DocxRGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt as PptPt


ROOT = Path(__file__).resolve().parents[1]
DELIVERABLES = ROOT / "deliverables"
DELIVERABLES.mkdir(exist_ok=True)

CV_TEMPLATE = ROOT / "APMP CV Template.docx"
PPIP_TEMPLATE = ROOT / "APMP Professional Impact Paper Template 1.2.pptx"

CV_OUTPUT = DELIVERABLES / "Timothy Semenza APMP Professional CV Draft.docx"
PPIP_OUTPUT = DELIVERABLES / "Timothy Semenza APMP Professional PPIP Draft.pptx"


CV_DATA = {
    "name": "Timothy Semenza",
    "contact": "Cinnaminson, NJ\n203-521-0311\ntimothy.semenza@gmail.com",
    "personal_statement": (
        "Proposal management leader and consultant with more than 11 years of experience across healthcare, "
        "insurance, federal, commercial, and facilities-services pursuits. Known for building governance where "
        "capture or proposal discipline is missing, leading cross-functional teams without formal authority, and "
        "converting high-pressure bids into reusable operating models. Current work through Boss Key LLC focuses "
        "on business-development acceleration, workflow improvement, and AI-enabled proposal operations."
    ),
    "experience": [
        "Boss Key LLC - Founder, Business Development Acceleration and Executive Coaching - February 2026 to Present\n"
        "Lead operational improvement engagements for service businesses focused on speed, quality, and profitability.\n"
        "Build practical workflow and automation solutions tied to measurable delivery improvement.\n"
        "Develop AI-enabled proposal and capture operating models for consulting clients.",
        "The Facilities Group - Senior Manager of Proposals - April 2024 to February 2026\n"
        "Designed and enforced enterprise proposal operating models across a multi-brand portfolio.\n"
        "Led the Detroit Public Schools recompete for RNA, coordinating about 15 stakeholders, reaching 2 finalist presentations, "
        "and creating content reused on about 12 later school bids.\n"
        "Standardized decision authority, compressed cycle time, and improved executive confidence in the proposal function.",
        "SBM Management Services, LP - Proposal Manager - October 2021 to April 2024\n"
        "Produced more than 200 compliant proposals in collaboration with sales, finance, and operations teams.\n"
        "Implemented a personal quality-control approach that reduced the need for post-submission clarification.\n"
        "Mentored newer team members and reinforced proposal-development standards.",
        "Kellermeyer Bergensons Services, LLC - Senior Proposal Coordinator - October 2020 to September 2021\n"
        "Guided proposal development from kickoff through submission with sales, finance, and operations stakeholders.\n"
        "Built a marketing content repository that improved reuse and bid consistency.\n"
        "Strengthened stakeholder partnership across business-development and delivery teams.",
        "Broadspire - RFP Developer - April 2019 to August 2020\n"
        "Managed healthcare-related RFP responses for third-party administrator services on pursuits over $1M.\n"
        "Developed standard language and branding resources in Qvidian.\n"
        "Provided consultative proposal support to sales teams.",
        "Blackstone Federal - Proposal Specialist - May 2018 to April 2019\n"
        "Supported federal proposal efforts including a Department of Homeland Security BPA pursuit valued up to $265M.\n"
        "Managed Shipley color-team reviews, compliance matrices, annotated outlines, and SME coordination.",
        "Deloitte Consulting - Proposal Writer - September 2017 to May 2018\n"
        "Led proposal development for private-sector opportunities valued at $25M+.\n"
        "Drafted and coordinated pursuit content using vetted firm information and repository assets.",
        "Broadspire - RFP Developer - April 2016 to September 2017\n"
        "Helped launch a new disability-insurance product category while drafting RFP responses.\n"
        "Built a proposal content library with more than 1,600 entries to improve efficiency and reuse.",
        "Aetna, a CVS Health Company - Proposal Writer - September 8, 2014 to April 2016\n"
        "Began proposal-management career supporting life and disability group insurance proposals.\n"
        "Worked with project managers and operations teams to tailor responses and strengthen sales strategy.",
    ],
    "publications": "Thought-leadership articles on government capture and AI-enabled proposal operations are planned for the Boss Key LLC website.",
    "qualifications": [
        "APMP Leadership Academy, issued May 2025",
        "APMP Micro-Certification: Bid and Proposal Writing, issued May 2023, Credential ID 2001186429",
        "APMP Capture Practitioner, issued February 2023, Credential ID 2001162940",
        "APMP Micro-Certification: Executive Summaries, issued January 2023, Credential ID 2001153927",
        "APMP Bid and Proposal Management Practitioner, issued July 2021, Credential ID 2001041238",
        "Leadership Challenge Program, issued June 2022",
        "ITIL Foundation Certificate in IT Service Management, issued September 2019",
    ],
    "education": "University of Connecticut, BA in English",
    "skills": (
        "Bid management; Proposal strategy; Executive summaries; Review management; Workflow design; "
        "Requirements gathering; Content management; Information gathering and analysis; Cross-functional team leadership; "
        "Continuous process improvement"
    ),
    "awards": "Breakout Room Hero recognition, Leadership Challenge Program",
    "references": (
        "Ann Couzins, Senior Manager of Pricing, The Facilities Group\n"
        "anncouzins@gmail.com | 513-277-1821\n"
        "Reference submitted separately through the APMP Professional process; final confirmation pending."
    ),
    "about_me": (
        "I am most effective in complex proposal environments where teams need structure, clear communication, and fast "
        "translation of live business input into compliant client-facing material. My background in English, proposal "
        "strategy, and workflow design supports a facilitative leadership style that balances quality, speed, and stakeholder confidence."
    ),
}


SLIDES = [
    {
        "title": "Proposal Professional Impact Paper",
        "subtitle": "Detroit/RNA Recompete Bid Rescue and Reuse",
        "body": ["Candidate Name: Timothy Semenza", "APMP Candidate Number: 28366", "Timeframe: November 2025 to January 2026"],
    },
    {
        "title": "Summary of Impact",
        "intro": (
            "In 20 working days, I took control of an unprepared $85M Detroit Public Schools recompete, coordinated about 15 stakeholders, reached 2 finalist presentations, and turned the work into reusable RNA school-bid infrastructure."
        ),
        "headers": ["Category", "Evidence"],
        "rows": [
            ["Strategic importance", "Incumbent work represented more than 25% of RNA's book of business and expanded from about $14M current scope to an $85M target"],
            ["Leadership role", "I led 100% of coordination and about 60% of drafting across executives, pricing, operations, engineering, Cheryl, and an external consultant"],
            ["Immediate outcome", "Compliant submission completed inside about 20 working days and advanced to 2 finalist presentations"],
            ["Sustained outcome", "Detroit-developed content and templates were reused on about 12 later school bids, including about 6 active at once immediately afterward"],
        ],
        "caption": "A high-risk recompete became a reusable school-bid operating model within one compressed pursuit cycle.",
    },
    {
        "title": "The Situation",
        "intro": (
            "RNA faced a Detroit Public Schools recompete with no meaningful capture plan, limited executive RFP engagement, "
            "and a late-added external consultant, even though the incumbent work represented more than 25% of RNA's book of business."
        ),
        "headers": ["Stakeholder Group", "Who Was Involved", "Why They Mattered"],
        "rows": [
            ["Senior leaders", "Victoria Erenze, Moufid Farah, Mohammed Mulf, Brandt Miller, Dave Angel", "Set direction, made decisions, and represented RNA in finalist activity"],
            ["Core delivery team", "Timothy Semenza, Chris Arlin", "Coordinated the pursuit and drafted the response"],
            ["Solution contributors", "Ann Couzins, Cheryl, Holly Stant, finance, operations, IT", "Supplied pricing, staffing, operational, and engineering inputs needed for a credible response"],
        ],
        "caption": "Proposal governance had to replace missing capture discipline before the team could compete credibly.",
    },
    {
        "title": "Stakeholder Engagement and Communication",
        "intro": "Direct engagement with senior and additional stakeholders was critical because the organization had limited preparation and I had to lead without formal authority.",
        "headers": ["Stakeholder Group", "How I Engaged Them", "Evidence of Leadership"],
        "rows": [
            ["Senior leaders", "Frequent working sessions, at peak as much as twice daily, plus live decision and review sessions", "I set direction, aligned next actions, and led the Tampa finalist strategy discussion under pressure"],
            ["Pricing and solution contributors", "Live interviews, pricing and staffing mediation, compliance tracking, and focused review of bid-specific changes", "I translated pricing and operational discussion into compliant narrative and kept disputed issues from stalling the bid"],
            ["Consultant and support team", "Shared company context, managed drafting split, and maintained SharePoint submission structure", "I remained the single point of coordination while enabling useful support from Chris Arlin"],
        ],
        "caption": "Direct engagement and communication created alignment where authority alone would not.",
    },
    {
        "title": "Tasks",
        "intro": "I stabilized the pursuit through five leadership tasks that gave the team structure, accountability, and decision discipline.",
        "headers": ["Task", "Why It Was Necessary", "Main Stakeholders Impacted"],
        "rows": [
            ["Establish bid governance and schedule", "No usable capture plan or agreed review cadence existed at release", "All senior stakeholders and the external consultant"],
            ["Develop evaluator and win-theme strategy", "The team needed a customer-centered strategy instead of an internally driven narrative", "RNA executives, Cheryl, and presentation speakers"],
            ["Elicit and structure SME input", "Live operational knowledge had to be turned into compliant proposal content", "Pricing, operations, engineering, and Cheryl"],
            ["Draft and quality-control content", "The written response and finalist materials had to stay aligned and submission-ready", "Consultant, executives, reviewers, and evaluators"],
            ["Lead review and submission readiness", "Readiness could not be left to last-minute judgment", "Executive reviewers, consultant, and submission support"],
        ],
        "caption": "Clear task ownership turned a fragmented effort into a manageable response.",
    },
    {
        "title": "Activities",
        "intro": "I used a milestone-led operating cadence to keep the pursuit on track from RFP receipt through two finalist rounds.",
        "headers": ["Checkpoint", "Approx Timing", "What Had To Be Complete", "How I Controlled It"],
        "rows": [
            ["RFP receipt and planning", "Day 1", "Automated schedule, project plan, and working cadence issued", "Power Automate schedule logic and immediate visibility of milestones"],
            ["Kickoff and evaluator strategy", "Days 1 to 2", "Leadership aligned on approach, evaluator research framed, responsibilities assigned", "Live kickoff plus customer-focused positioning before drafting expanded"],
            ["Compliance matrix and SME baseline", "Days 2 to 5", "Compliance matrix built and core pricing, operations, and account insight captured", "Interviews, Teams follow-up, and tracked open items"],
            ["Draft build and review cycle", "Days 6 to 12", "Sections drafted, consultant integrated, QA performed, and pink/red-style reviews completed", "Read-aloud QA, narrative control, and recurring working sessions"],
            ["Executive review and submission", "Days 13 to 20", "Executive review package, final files, SharePoint structure, and compliant submission ready", "Readiness checks, clear folder structure, and planned handoff coverage"],
            ["Finalist rounds 1 and 2", "Post-submission", "Speaker notes, slide messaging, likely questions, and live strategy alignment completed", "Tampa war room, action plans, and repeated coordination under short notice"],
        ],
        "bullets": [
            "Cadence peaked at twice daily with senior leaders, while AI and automation accelerated planning, compliance extraction, and draft production inside a candidate-led process.",
        ],
        "caption": "Automated planning and recurring governance made a short public-sector timeline executable.",
    },
    {
        "title": "Results",
        "intro": "I measured the impact through finalist progression, reuse, draft speed, and executive review effort rather than through final award visibility.",
        "headers": ["Metric", "Baseline", "Result", "How Measured"],
        "rows": [
            ["Finalist progression", "No assured advancement from an unprepared start", "The bid reached 2 finalist presentations", "Invitation to two finalist rounds"],
            ["Reusable bid base", "Each school bid rebuilt structure and messaging", "About 12 later RNA school bids reused Detroit assets, with about 6 active at once immediately afterward", "Observed live reuse across later school-bid production"],
            ["First-draft speed", "About 4 hours to reach a first compliant draft", "An about 80% draft could be reached in roughly 5 minutes before tailoring", "Compared live prep effort on later similar school bids"],
            ["Executive review burden", "Leaders saw unfamiliar structure and needed repeated writing review", "Later bids usually needed only 1 executive review focused on solution and pricing differences", "Observed post-Detroit review pattern with executives and pricing"],
        ],
        "bullets": [
            "Later school bids shifted toward about 3 hours of pricing and solutioning meetings plus light administrative completion rather than rebuilding the narrative base.",
            "Final award result is unknown, so the claim is anchored in finalist progression, reuse, speed, review effort, and Ann-corroborated turnaround patterns.",
        ],
        "caption": "Even without final award visibility, the pursuit produced measurable reuse, speed, credibility, and submission outcomes.",
    },
    {
        "title": "CPD Plan for the Next 24 Months",
        "intro": "My next 24 months of development are focused on deepening capture capability, formalizing AI-enabled proposal delivery, and contributing more visibly to the profession.",
        "headers": ["Timing", "Activity", "Why It Matters"],
        "rows": [
            ["March to April 2026", "Finish and live-smoke-test my internally developed capture-to-proposal application", "Validate a purpose-built operating system for capture and proposal work"],
            ["April to May 2026", "Actively build government-capture capability", "Expand facilities-services advisory work further upstream in the selling cycle"],
            ["March to June 2026", "Earn the APMP AI Micro-Certification", "Align AI-enabled practice with the official APMP credential"],
            ["March to September 2026", "Increase involvement in the APMP National Capital Area chapter", "Contribute more directly through events, speaking, and board-oriented service"],
            ["By March 2027", "Stand up a turnkey AI-enabled consulting model and develop toward a BPC 2027 webinar", "Demonstrate a practical end-to-end AI operating model for capture and proposal management"],
        ],
        "caption": "The CPD plan is tied directly to service expansion, standards alignment, and visible contribution to the profession.",
    },
    {
        "title": "CPD Application and Knowledge Sharing",
        "intro": "I intend to convert learning into reusable capability for clients and into practical advocacy for the proposal profession.",
        "headers": ["Learning Area", "How I Will Apply It", "How I Will Share It"],
        "rows": [
            ["Government capture", "Move clients upstream from post-RFP response into earlier capture-stage planning", "Publish insight articles on the Boss Key LLC website"],
            ["AI-enabled proposal operations", "Use repo-based conversational AI with human judgment in the loop to accelerate compliant delivery", "Share methods with APMP chapter peers and through online presentations"],
            ["Community leadership", "Increase contribution to the NCA APMP chapter through events and speaking", "Support chapter events and develop conference-level content"],
            ["Standards-led innovation", "Anchor experimentation in recognized APMP good practice", "Translate lessons from the APMP AI Micro-Certification into client delivery and peer discussions"],
        ],
        "bullets": [
            "The same principle underpins both my impact story and my development plan: the best proposal leadership leaves behind a stronger system, not just a stronger single submission."
        ],
        "caption": "Learning is being converted into reusable client capability and visible contribution to the proposal profession.",
    },
]

INTRO_TOP = Inches(1.48)
INTRO_HEIGHT = Inches(0.5)
TABLE_TOP = Inches(2.0)
TABLE_HEIGHT_FULL = Inches(4.1)
TABLE_HEIGHT_COMPACT = Inches(3.2)
CAPTION_TOP = Inches(6.82)
INTRO_WIDTH = Inches(10.25)


def remove_shape(shape):
    element = shape._element
    element.getparent().remove(element)


def add_textbox(slide, left, top, width, height, text, size=16, bold=False, italic=False, color=(0, 0, 0)):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = text
    run.font.size = PptPt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(*color)
    return box


def style_table(table, header_size=12, body_size=11):
    for j in range(len(table.columns)):
        cell = table.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(31, 78, 121)
        for p in cell.text_frame.paragraphs:
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = PptPt(header_size)
    for i in range(1, len(table.rows)):
        for j in range(len(table.columns)):
            cell = table.cell(i, j)
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = PptPt(body_size)
                    run.font.color.rgb = RGBColor(0, 0, 0)


def add_caption(slide, text):
    add_textbox(slide, Inches(0.7), CAPTION_TOP, Inches(11.7), Inches(0.28), text, size=11, italic=True, color=(89, 89, 89))


def reset_slide(slide, title):
    has_title = False
    for shape in list(slide.shapes):
        if shape.is_placeholder and shape.placeholder_format.type == 1:
            shape.text = title
            has_title = True
        else:
            remove_shape(shape)
    if not has_title:
        add_textbox(slide, Inches(0.55), Inches(0.2), Inches(12.1), Inches(0.45), title, size=22, bold=True)


def build_ppip():
    prs = Presentation(str(PPIP_TEMPLATE))

    cover = prs.slides[0]
    subtitle_shape = cover.shapes[0]
    title_shape = cover.shapes[1]
    body_shape = cover.shapes[2]

    subtitle_shape.text = ""

    title_shape.text = SLIDES[0]["title"]
    title_shape.left = Inches(5.55)
    title_shape.top = Inches(1.48)
    title_shape.width = Inches(6.3)
    title_shape.height = Inches(0.6)
    for p in title_shape.text_frame.paragraphs:
        for run in p.runs:
            run.font.size = PptPt(24)
            run.font.color.rgb = RGBColor(102, 102, 102)

    body_shape.text = ""
    add_textbox(cover, Inches(5.58), Inches(2.18), Inches(6.0), Inches(0.8), SLIDES[0]["subtitle"], size=18, color=(31, 78, 121))
    add_textbox(cover, Inches(5.6), Inches(3.45), Inches(5.4), Inches(1.4), "\n".join(SLIDES[0]["body"]), size=14, color=(31, 31, 31))

    for i, content in enumerate(SLIDES[1:], start=1):
        slide = prs.slides[i]
        reset_slide(slide, content["title"])
        add_textbox(slide, Inches(0.78), INTRO_TOP, INTRO_WIDTH, INTRO_HEIGHT, content["intro"], size=14, color=(31, 31, 31))
        table_h = TABLE_HEIGHT_FULL if i in (1, 2, 3, 4, 7) else TABLE_HEIGHT_COMPACT
        header_size = 12
        body_size = 11
        if i == 5:
            table_h = Inches(3.92)
            header_size = 11
            body_size = 10
        elif i == 6:
            table_h = Inches(3.8)
            header_size = 11
            body_size = 10.2
        table = slide.shapes.add_table(len(content["rows"]) + 1, len(content["headers"]), Inches(0.7), TABLE_TOP, Inches(12.0), table_h).table
        for j, header in enumerate(content["headers"]):
            table.cell(0, j).text = header
        for r, row in enumerate(content["rows"], start=1):
            for c, val in enumerate(row):
                table.cell(r, c).text = val
        style_table(table, header_size=header_size, body_size=body_size)
        if "bullets" in content:
            y = Inches(6.02) if i == 5 else Inches(5.92) if i == 6 else Inches(5.38) if i != 8 else Inches(5.58)
            bullet_height = Inches(0.48) if i == 5 else Inches(0.62) if i == 6 else Inches(0.92)
            bullet_size = 10.5 if i == 5 else 10.8 if i == 6 else 11.5
            bullet_box = slide.shapes.add_textbox(Inches(0.92), y, Inches(11.1), bullet_height)
            tf = bullet_box.text_frame
            tf.word_wrap = True
            for idx, bullet in enumerate(content["bullets"]):
                p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
                p.text = bullet
                p.font.size = PptPt(bullet_size)
                p.font.color.rgb = RGBColor(31, 54, 99)
                p.bullet = True
        add_caption(slide, content["caption"])

    prs.save(str(PPIP_OUTPUT))


def build_cv():
    doc = Document()

    section = doc.sections[0]
    section.top_margin = DocxInches(0.55)
    section.bottom_margin = DocxInches(0.55)
    section.left_margin = DocxInches(0.65)
    section.right_margin = DocxInches(0.65)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)

    def add_section_heading(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(text.upper())
        run.font.name = "Aptos"
        run.font.size = Pt(10.5)
        run.font.bold = True
        run.font.color.rgb = DocxRGBColor(31, 78, 121)

    def add_body(text, before=0, after=4):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.line_spacing = 1.08
        run = p.add_run(text)
        run.font.name = "Aptos"
        run.font.size = Pt(10)
        run.font.color.rgb = DocxRGBColor(31, 31, 31)
        return p

    def add_bullet(text):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.04
        p.text = text
        for run in p.runs:
            run.font.name = "Aptos"
            run.font.size = Pt(9.5)
            run.font.color.rgb = DocxRGBColor(31, 31, 31)

    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.paragraph_format.space_after = Pt(1)
    run = name.add_run(CV_DATA["name"])
    run.font.name = "Aptos Display"
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = DocxRGBColor(31, 78, 121)

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(8)
    c_run = contact.add_run("Cinnaminson, NJ | 203-521-0311 | timothy.semenza@gmail.com")
    c_run.font.name = "Aptos"
    c_run.font.size = Pt(9.5)
    c_run.font.color.rgb = DocxRGBColor(89, 89, 89)

    add_section_heading("Personal Statement")
    add_body(CV_DATA["personal_statement"])

    add_section_heading("Professional Experience")
    for block in CV_DATA["experience"]:
        lines = block.split("\n")
        role = doc.add_paragraph()
        role.paragraph_format.space_before = Pt(6)
        role.paragraph_format.space_after = Pt(2)
        r = role.add_run(lines[0])
        r.font.name = "Aptos"
        r.font.size = Pt(10.5)
        r.font.bold = True
        r.font.color.rgb = DocxRGBColor(31, 31, 31)
        for bullet in lines[1:]:
            add_bullet(bullet)

    add_section_heading("Publications")
    add_body(CV_DATA["publications"])

    add_section_heading("Qualifications")
    for item in CV_DATA["qualifications"]:
        add_bullet(item)

    add_section_heading("Education")
    add_body(CV_DATA["education"], after=2)

    add_section_heading("Skills")
    add_body(CV_DATA["skills"], after=2)

    add_section_heading("Awards")
    add_body(CV_DATA["awards"], after=2)

    add_section_heading("References")
    add_body(CV_DATA["references"], after=2)

    add_section_heading("About Me")
    add_body(CV_DATA["about_me"])

    doc.save(str(CV_OUTPUT))


def main():
    build_cv()
    build_ppip()
    print(f"Wrote {CV_OUTPUT}")
    print(f"Wrote {PPIP_OUTPUT}")


if __name__ == "__main__":
    main()
