from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from pathlib import Path

wb = Workbook()

# Styles
header_fill = PatternFill('solid', fgColor='0B1020')
header_font = Font(color='FFFFFF', bold=True)
sub_fill = PatternFill('solid', fgColor='E8EFFF')
sub_font = Font(color='0B1020', bold=True)
label_font = Font(bold=True, color='0B1020')
thin = Side(style='thin', color='D7DEEB')
border = Border(left=thin, right=thin, top=thin, bottom=thin)

# Inputs sheet
ws = wb.active
ws.title = 'Inputs'
ws['A1'] = 'Boss Key Financial Architecture Model'
ws['A1'].font = Font(size=14, bold=True, color='0B1020')
ws.merge_cells('A1:E1')

ws['A2'] = 'Scenario Selection'
ws['B2'] = 'Moderate'
ws['D2'] = 'Use dropdown in B2'
ws['D2'].font = Font(italic=True, color='505A71')

dv = DataValidation(type='list', formula1='"Conservative,Moderate"', allow_blank=False)
ws.add_data_validation(dv)
dv.add(ws['B2'])

ws['A4'] = 'Assumption'
ws['B4'] = 'Conservative'
ws['C4'] = 'Moderate'
ws['D4'] = 'Selected'
for c in 'ABCD':
    cell = ws[f'{c}4']
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center')
    cell.border = border

assumptions = [
    ('Business Retainer Clients (avg)', 4.0, 4.0),
    ('Seller Retainer Clients (avg)', 0.5, 1.0),
    ('Avg Business Retainer ($/mo)', 5600, 5750),
    ('Avg Seller Retainer ($/mo)', 2400, 2500),
    ('Business Diagnostics per Year', 6, 6),
    ('Avg Business Diagnostic Fee ($)', 5200, 5500),
    ('Seller Diagnostics per Year', 3, 4),
    ('Avg Seller Diagnostic Fee ($)', 1600, 1750),
    ('Non-Owner Overhead ($/yr)', 45000, 45000),
    ('Reasonable Salary ($/yr)', 130000, 140000),
    ('Employer Payroll Tax Rate', 0.0765, 0.0765),
    ('Federal Effective Tax Rate', 0.21, 0.23),
    ('State Effective Tax Rate', 0.06, 0.065),
    ('Quarterly Tax Reserve % (of distributions)', 0.38, 0.38),
    ('Operating Buffer Target % of Revenue', 0.15, 0.15),
]

start = 5
for i, (name, cons, mod) in enumerate(assumptions, start=start):
    ws[f'A{i}'] = name
    ws[f'B{i}'] = cons
    ws[f'C{i}'] = mod
    ws[f'D{i}'] = f'=IF($B$2="Conservative",B{i},C{i})'
    for c in 'ABCD':
        ws[f'{c}{i}'].border = border

for r in range(5, 20):
    if r in (15,16,17,18,19):
        ws[f'B{r}'].number_format = '0.00%'
        ws[f'C{r}'].number_format = '0.00%'
        ws[f'D{r}'].number_format = '0.00%'

for r in [7,8,10,12,13,14]:
    ws[f'B{r}'].number_format = '$#,##0'
    ws[f'C{r}'].number_format = '$#,##0'
    ws[f'D{r}'].number_format = '$#,##0'

ws['A22'] = 'Target Annual Revenue'
ws['B22'] = 350000
ws['A23'] = 'Prior W-2 Salary Reference'
ws['B23'] = 164000
for r in [22,23]:
    ws[f'A{r}'].font = label_font
    ws[f'B{r}'].number_format = '$#,##0'

# Annual Model sheet
am = wb.create_sheet('Annual_Model')
am['A1'] = 'Annual Financial Architecture'
am['A1'].font = Font(size=14, bold=True, color='0B1020')
am.merge_cells('A1:D1')

for c, h in enumerate(['Metric', 'Formula/Driver', 'Amount', 'Notes'], start=1):
    cell = am.cell(row=3, column=c, value=h)
    cell.fill = header_fill
    cell.font = header_font
    cell.border = border

rows = [
    ('Retainer Revenue', '=(Inputs!D5*Inputs!D7 + Inputs!D6*Inputs!D8)*12', 'Business + seller retainers'),
    ('Diagnostic Revenue', '=Inputs!D9*Inputs!D10 + Inputs!D11*Inputs!D12', 'Business + seller diagnostics'),
    ('Gross Revenue', '=C4+C5', ''),
    ('Non-Owner Overhead', '=Inputs!D13', ''),
    ('Pre-Owner Comp Operating Profit', '=C6-C7', ''),
    ('Reasonable Salary', '=Inputs!D14', ''),
    ('Employer Payroll Tax', '=C9*Inputs!D15', ''),
    ('S-Corp Profit (Pre-Income Tax)', '=C8-C9-C10', ''),
    ('Combined Effective Tax Rate', '=Inputs!D16+Inputs!D17', ''),
    ('Estimated Income Tax on Salary + Profit', '=(C9+C11)*C12', 'Planning estimate'),
    ('Estimated After-Tax Owner Cash', '=C9+C11-C13', 'Salary + distributions net of est tax'),
    ('Operating Buffer Target', '=C6*Inputs!D19', ''),
]

r = 4
for metric, formula, note in rows:
    am[f'A{r}'] = metric
    am[f'B{r}'] = formula
    am[f'C{r}'] = f'={formula}' if not formula.startswith('=') else formula
    am[f'D{r}'] = note
    for c in 'ABCD':
        am[f'{c}{r}'].border = border
    r += 1

for row in range(4, r):
    am[f'C{row}'].number_format = '$#,##0'
am['C12'].number_format = '0.00%'

for highlight in [6,8,11,14,15]:
    for c in 'ABCD':
        am[f'{c}{highlight}'].fill = sub_fill
        am[f'{c}{highlight}'].font = sub_font

# Monthly Cashflow sheet
mc = wb.create_sheet('Monthly_Cashflow')
mc['A1'] = 'Monthly Cash Flow Discipline (Planning)'
mc['A1'].font = Font(size=14, bold=True, color='0B1020')
mc.merge_cells('A1:J1')

headers = ['Month', 'Revenue', 'Overhead', 'Salary', 'Employer Payroll Tax', 'Pre-Tax Profit', 'Tax Reserve Transfer', 'Owner Distribution (Gross)', 'Ending Operating Buffer', 'Notes']
for i, h in enumerate(headers, start=1):
    cell = mc.cell(row=3, column=i, value=h)
    cell.fill = header_fill
    cell.font = header_font
    cell.border = border

months = ['Jan','Feb','Mar','Apr','May','Jun','Jul','Aug','Sep','Oct','Nov','Dec']
for idx, m in enumerate(months, start=4):
    mc[f'A{idx}'] = m
    mc[f'B{idx}'] = '=Annual_Model!C6/12'
    mc[f'C{idx}'] = '=Inputs!D13/12'
    mc[f'D{idx}'] = '=Inputs!D14/12'
    mc[f'E{idx}'] = '=D{0}*Inputs!D15'.format(idx)
    mc[f'F{idx}'] = '=B{0}-C{0}-D{0}-E{0}'.format(idx)
    mc[f'G{idx}'] = '=MAX(0,F{0}*Inputs!D18)'.format(idx)
    mc[f'H{idx}'] = '=MAX(0,F{0}-G{0})'.format(idx)
    if idx == 4:
        mc[f'I{idx}'] = '=Inputs!B22*Inputs!D19/2 + H4'
    else:
        mc[f'I{idx}'] = '=I{0}+H{1}'.format(idx-1, idx)
    mc[f'J{idx}'] = ''
    for c in range(1, 11):
        mc.cell(row=idx, column=c).border = border

for row in range(4, 16):
    for col in ['B','C','D','E','F','G','H','I']:
        mc[f'{col}{row}'].number_format = '$#,##0'

mc['A17'] = 'Annual Totals'
mc['B17'] = '=SUM(B4:B15)'
mc['C17'] = '=SUM(C4:C15)'
mc['D17'] = '=SUM(D4:D15)'
mc['E17'] = '=SUM(E4:E15)'
mc['F17'] = '=SUM(F4:F15)'
mc['G17'] = '=SUM(G4:G15)'
mc['H17'] = '=SUM(H4:H15)'
mc['I17'] = '=I15'
for c in 'ABCDEFGHI':
    mc[f'{c}17'].fill = sub_fill
    mc[f'{c}17'].font = sub_font
    mc[f'{c}17'].border = border

for col in ['B','C','D','E','F','G','H','I']:
    mc[f'{col}17'].number_format = '$#,##0'

# Scenario compare
sc = wb.create_sheet('Scenario_Compare')
sc['A1'] = 'Scenario Comparison (Conservative / Moderate / Strong)'
sc['A1'].font = Font(size=14, bold=True, color='0B1020')
sc.merge_cells('A1:G1')

for i, h in enumerate(['Metric', 'Conservative', 'Moderate', 'Strong', 'Notes'], start=1):
    cell = sc.cell(row=3, column=i, value=h)
    cell.fill = header_fill
    cell.font = header_font
    cell.border = border

scenarios = [
    ('Business Clients', '=Inputs!B5', '=Inputs!C5', 5.0, ''),
    ('Seller Clients', '=Inputs!B6', '=Inputs!C6', 1.0, ''),
    ('Avg Business Retainer ($/mo)', '=Inputs!B7', '=Inputs!C7', 6000, ''),
    ('Avg Seller Retainer ($/mo)', '=Inputs!B8', '=Inputs!C8', 2800, ''),
    ('Retainer Revenue', '=(B4*B6 + B5*B7)*12', '=(C4*C6 + C5*C7)*12', '=(D4*D6 + D5*D7)*12', ''),
    ('Business Diagnostics (#)', '=Inputs!B9', '=Inputs!C9', 7, ''),
    ('Avg Business Diagnostic Fee', '=Inputs!B10', '=Inputs!C10', 6000, ''),
    ('Seller Diagnostics (#)', '=Inputs!B11', '=Inputs!C11', 4, ''),
    ('Avg Seller Diagnostic Fee', '=Inputs!B12', '=Inputs!C12', 2000, ''),
    ('Diagnostic Revenue', '=B9*B10 + B11*B12', '=C9*C10 + C11*C12', '=D9*D10 + D11*D12', ''),
    ('Total Revenue', '=B8+B13', '=C8+C13', '=D8+D13', ''),
]

row = 4
for metric, b, c, d, note in scenarios:
    sc[f'A{row}'] = metric
    sc[f'B{row}'] = b
    sc[f'C{row}'] = c
    sc[f'D{row}'] = d
    sc[f'E{row}'] = note
    for col in 'ABCDE':
        sc[f'{col}{row}'].border = border
    row += 1

for r in [6,7,8,10,12,13,14]:
    for col in ['B','C','D']:
        sc[f'{col}{r}'].number_format = '$#,##0'

# Notes sheet
nt = wb.create_sheet('Notes')
nt['A1'] = 'Model Notes'
nt['A1'].font = Font(size=14, bold=True, color='0B1020')
notes = [
    '1) This is a planning model, not tax/legal advice.',
    '2) Change Inputs!B2 between Conservative and Moderate to toggle the model.',
    '3) Salary level should be reviewed with CPA for reasonable compensation support.',
    '4) Tax reserve strategy assumes disciplined quarterly set-asides.',
    '5) Keep client roster capped to preserve delivery quality and margin.',
    '6) Roster mix assumptions separate business and seller client economics.',
    '7) Moderate mix target: 4 business retainers + 1 seller retainer, with 6 business and 4 seller diagnostics/year.',
]
for i, line in enumerate(notes, start=3):
    nt[f'A{i}'] = line

# Column widths
for sheet in [ws, am, mc, sc, nt]:
    widths = {
        'A': 42, 'B': 24, 'C': 24, 'D': 28, 'E': 26, 'F': 20, 'G': 24, 'H': 24, 'I': 24, 'J': 28
    }
    for col, w in widths.items():
        sheet.column_dimensions[col].width = w

out_path = Path(__file__).resolve().with_name("Boss-Key-Financial-Architecture.xlsx")
wb.save(out_path)
print(out_path)
