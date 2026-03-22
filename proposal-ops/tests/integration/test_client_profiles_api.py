from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient

from app.modules.proposal_builder.service import ProposalBuilderService


class _UnavailableOpenAIClient:
    available = False
    model = "gpt-5.4"
    timeout_seconds = 300


def _proposal_template_bytes() -> bytes:
    document = Document()
    document.add_paragraph("{{ opportunity_name }}")
    document.add_paragraph("{{ transmittal_letter }}")
    document.add_paragraph("{{ executive_summary }}")
    document.add_paragraph("{{ technical_approach }}")
    document.add_paragraph("{{ management_staffing_plan }}")
    document.add_paragraph("{{ transition_mobilization_plan }}")
    document.add_paragraph("{{ past_performance }}")
    document.add_paragraph("{{ pricing_and_commercials }}")
    document.add_paragraph("{{ compliance_attachments }}")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_client_profile_api_supports_profile_defaults_and_builder_selection(client: TestClient, monkeypatch) -> None:
    monkeypatch.setattr(ProposalBuilderService, "_openai_client", lambda self, **_kwargs: _UnavailableOpenAIClient())

    created_profile = client.post(
        "/api/client-profiles",
        json={
            "display_name": "Metro Regional Airport Authority",
            "aliases": ["Metro Airport Authority"],
            "approved_content_tags": ["airport", "janitorial"],
            "region_tags": ["northeast"],
            "service_tags": ["janitorial", "day-porter"],
        },
    )
    assert created_profile.status_code == 200
    profile = created_profile.json()
    profile_id = profile["id"]

    created_pricing_model = client.post(
        f"/api/client-profiles/{profile_id}/pricing-models",
        json={
            "name": "Airport Janitorial Pricing Model",
            "assumptions": ["Use approved airport staffing assumptions."],
            "service_tags": ["janitorial", "day-porter"],
            "region_tags": ["northeast"],
            "site_multiplier": 1.0,
            "continuous_coverage_multiplier": 1.35,
            "labor_categories": [
                {
                    "code": "janitor",
                    "label": "Janitor",
                    "hourly_rate": 20.5,
                    "burden_factor": 1.27,
                    "markup_factor": 1.18,
                    "default_hours_per_week": 160,
                },
                {
                    "code": "day_porter",
                    "label": "Day Porter",
                    "hourly_rate": 22.0,
                    "burden_factor": 1.24,
                    "markup_factor": 1.16,
                    "default_hours_per_week": 80,
                },
            ],
        },
    )
    assert created_pricing_model.status_code == 200
    pricing_model = created_pricing_model.json()

    template_upload = client.post(
        f"/api/client-profiles/{profile_id}/proposal-templates",
        data={"name": "Authority Proposal Template"},
        files={
            "file": (
                "authority-proposal-template.docx",
                _proposal_template_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert template_upload.status_code == 200
    proposal_template = template_upload.json()
    assert proposal_template["status"] == "APPROVED"

    defaults = client.post(
        f"/api/client-profiles/{profile_id}/defaults",
        json={
            "default_pricing_model_id": pricing_model["id"],
            "default_proposal_template_id": proposal_template["id"],
        },
    )
    assert defaults.status_code == 200
    assert defaults.json()["default_pricing_model_id"] == pricing_model["id"]
    assert defaults.json()["default_proposal_template_id"] == proposal_template["id"]

    started = client.post(
        "/proposal-builder/start",
        data={"use_demo_sample": "true", "actor": "operator"},
        follow_redirects=False,
    )
    assert started.status_code == 303
    opportunity_id = started.headers["location"].split("/")[2]

    workspace = client.get(f"/api/opportunities/{opportunity_id}/proposal-builder")
    assert workspace.status_code == 200
    workspace_payload = workspace.json()
    assert workspace_payload["proposal_config"]["client_profile_name"] == "Metro Regional Airport Authority"
    assert workspace_payload["proposal_config"]["pricing_model_name"] == "Airport Janitorial Pricing Model"
    assert workspace_payload["proposal_config"]["proposal_template_name"] == "Authority Proposal Template"
    assert workspace_payload["proposal_config"]["client_profile_selection_source"] == "default"

    exported = client.post(
        f"/api/opportunities/{opportunity_id}/proposal-builder/export",
        json={"actor": "operator"},
    )
    assert exported.status_code == 200
    export_payload = exported.json()
    assert export_payload["proposal_config"]["proposal_template_name"] == "Authority Proposal Template"
    assert export_payload["export_manifest"]["proposal_template_name"] == "Authority Proposal Template"
