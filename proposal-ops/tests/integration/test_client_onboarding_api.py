from io import BytesIO
import json
import zipfile

from docx import Document
from fastapi.testclient import TestClient

from app.modules.proposal_builder.service import ProposalBuilderService


class _UnavailableOpenAIClient:
    available = False
    model = "gpt-5.4"
    timeout_seconds = 300


def _proposal_template_bytes() -> bytes:
    document = Document()
    document.add_paragraph("{{ opportunity_name }}")
    document.add_paragraph("{{ client_name }}")
    document.add_paragraph("{{ solicitation_number }}")
    document.add_paragraph("{{ proposal_due }}")
    document.add_paragraph("{{ transmittal_letter }}")
    document.add_paragraph("{{ executive_summary }}")
    document.add_paragraph("{{ technical_approach }}")
    document.add_paragraph("{{ management_staffing_plan }}")
    document.add_paragraph("{{ transition_mobilization_plan }}")
    document.add_paragraph("{{ past_performance }}")
    document.add_paragraph("{{ pricing_and_commercials }}")
    document.add_paragraph("{{ compliance_attachments }}")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _onboarding_zip_bytes() -> bytes:
    payload = BytesIO()
    with zipfile.ZipFile(payload, "w") as archive:
        archive.writestr(
            "client-onboarding.json",
            json.dumps(
                {
                    "client": {
                        "display_name": "Metro Regional Airport Authority",
                        "aliases": ["Metro Airport Authority"],
                        "approved_content_tags": ["airport", "janitorial"],
                        "region_tags": ["northeast"],
                        "service_tags": ["janitorial", "day-porter"],
                    },
                    "playbook": {
                        "default_tone": "formal, buyer-focused, and operationally precise",
                        "preferred_terminology": ["Authority", "contract manager", "site supervision"],
                        "avoid_terminology": ["cheap", "generic"],
                        "required_attachments": ["Pricing workbook", "Certification forms", "Insurance evidence"],
                    },
                }
            ),
        )
        archive.writestr(
            "pricing-model.csv",
            "labor category,hourly rate,burden factor,markup factor,default hours per week\n"
            "Janitor,20.50,1.27,1.18,160\n"
            "Day Porter,22.00,1.24,1.16,80\n",
        )
        archive.writestr("client-playbook.txt", "Emphasize reporting visibility, local supervision, and public-facing operations.")
        archive.writestr("resume-ops-director.txt", "Operations director resume content.")
        archive.writestr("reference-airport.txt", "Comparable airport janitorial reference.")
        archive.writestr("brand-logo.txt", "Brand asset placeholder.")
        archive.writestr("authority-proposal-template.docx", _proposal_template_bytes())
    return payload.getvalue()


def test_client_onboarding_api_uploads_approves_and_activates_pack(client: TestClient, monkeypatch) -> None:
    monkeypatch.setattr(ProposalBuilderService, "_openai_client", lambda self, **_kwargs: _UnavailableOpenAIClient())

    created = client.post(
        "/api/client-onboarding/packs",
        data={
            "actor": "admin",
            "pack_name": "Airport Authority Pack",
            "client_display_name": "Metro Regional Airport Authority",
        },
        files={"archive": ("airport-authority-pack.zip", _onboarding_zip_bytes(), "application/zip")},
    )
    assert created.status_code == 200
    pack = created.json()["pack"]
    assert pack["status"] == "VALIDATED"
    assert any(asset["asset_role"] == "proposal_template" for asset in pack["assets"])
    assert any(asset["asset_role"] == "pricing_artifact" for asset in pack["assets"])
    assert pack["playbook"]["title"] == "Metro Regional Airport Authority Proposal Playbook"

    approved = client.post(
        f"/api/client-onboarding/packs/{pack['id']}/approve",
        json={"actor": "admin"},
    )
    assert approved.status_code == 200
    assert approved.json()["status"] == "APPROVED"

    activated = client.post(
        f"/api/client-onboarding/packs/{pack['id']}/activate",
        json={"actor": "admin"},
    )
    assert activated.status_code == 200
    assert activated.json()["activated_at"] is not None

    started = client.post(
        "/proposal-builder/start",
        data={"use_demo_sample": "true", "actor": "operator"},
        follow_redirects=False,
    )
    assert started.status_code == 303
    opportunity_id = started.headers["location"].split("/")[2]

    workspace = client.get(f"/api/opportunities/{opportunity_id}/proposal-builder")
    assert workspace.status_code == 200
    payload = workspace.json()
    assert payload["active_onboarding_pack_id"] == pack["id"]
    assert payload["client_environment_status"] == "READY"
    assert payload["proposal_config"]["active_onboarding_pack_name"] == "Airport Authority Pack"
    assert payload["proposal_config"]["active_playbook_version"] == 1
    assert payload["proposal_config"]["client_profile_name"] == "Metro Regional Airport Authority"
    assert not payload["setup_gaps"]


def test_client_onboarding_api_supports_asset_override_and_playbook_revision(client: TestClient) -> None:
    created = client.post(
        "/api/client-onboarding/packs",
        data={
            "actor": "admin",
            "pack_name": "Editable Pack",
            "client_display_name": "Metro Regional Airport Authority",
        },
        files={"archive": ("airport-authority-pack.zip", _onboarding_zip_bytes(), "application/zip")},
    )
    assert created.status_code == 200
    pack = created.json()["pack"]

    brand_asset = next(asset for asset in pack["assets"] if asset["asset_role"] == "brand_asset")
    updated = client.post(
        f"/api/client-onboarding/packs/{pack['id']}/assets/{brand_asset['id']}",
        json={"actor": "admin", "asset_role": "approved_content", "asset_status": "CANDIDATE"},
    )
    assert updated.status_code == 200
    updated_pack = updated.json()
    edited_asset = next(asset for asset in updated_pack["assets"] if asset["id"] == brand_asset["id"])
    assert edited_asset["asset_role"] == "approved_content"
    assert edited_asset["classification_source"] == "manual"

    saved_playbook = client.post(
        f"/api/client-onboarding/packs/{pack['id']}/playbook",
        json={
            "actor": "admin",
            "title": "Metro Authority Proposal Playbook",
            "narrative_guidance_md": "Lead with buyer confidence, operational control, and airport-ready supervision.",
        },
    )
    assert saved_playbook.status_code == 200
    saved_pack = saved_playbook.json()
    assert saved_pack["playbook"]["title"] == "Metro Authority Proposal Playbook"
    assert "airport-ready supervision" in saved_pack["playbook"]["narrative_guidance_md"]

    approved = client.post(
        f"/api/client-onboarding/packs/{pack['id']}/approve",
        json={"actor": "admin"},
    )
    assert approved.status_code == 200

    revised = client.post(
        f"/api/client-onboarding/packs/{pack['id']}/playbook",
        json={
            "actor": "admin",
            "title": "Metro Authority Proposal Playbook",
            "narrative_guidance_md": "Refresh the message around accountability, reporting visibility, and responsive leadership.",
        },
    )
    assert revised.status_code == 200
    revised_pack = revised.json()
    assert revised_pack["status"] == "VALIDATED"
    assert revised_pack["playbook"]["version"] == 2
    assert "responsive leadership" in revised_pack["playbook"]["narrative_guidance_md"]
