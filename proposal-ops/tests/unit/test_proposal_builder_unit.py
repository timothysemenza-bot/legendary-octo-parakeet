from pathlib import Path

from docx import Document

from app.modules.proposal_builder.content_library import content_block_inventory, load_content_blocks
import app.modules.proposal_builder.export as export_module
from app.modules.proposal_builder.export import build_export_bundle
from app.modules.proposal_builder.service import _coerce_draft_stage_artifact
from app.modules.proposal_builder.template_renderer import validate_proposal_template


def test_content_library_inventory_loads_blocks() -> None:
    blocks = load_content_blocks()
    assert blocks
    inventory = content_block_inventory()
    assert inventory
    assert any(item["block_id"].startswith("sample-") for item in inventory)


def test_export_bundle_writes_docx_and_zip(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr(export_module, "EXPORTS_DIR", tmp_path)
    workspace = {
        "opportunity_id": "opp-demo",
        "opportunity_name": "Airport Terminal Janitorial and Day Porter Services",
        "client_name": "Metro Regional Airport Authority",
        "solicitation_id": "sol-demo",
        "extracted_deadline": "2026-04-10",
        "opportunity_summary": {
            "client_name": "Metro Regional Airport Authority",
            "opportunity_name": "Airport Terminal Janitorial and Day Porter Services",
            "solicitation_number": "MR-AA-2026-041",
            "proposal_due_date": "2026-04-10",
            "proposal_due_time": "2:00 PM ET",
            "strategic_fit": "Good",
            "geographic_fit": "Known",
            "operational_fit": "Known",
            "complexity_level": "Moderate",
            "risk_level": "Medium",
            "recommended_next_action": "Proceed",
        },
        "win_themes": [{"title": "Low-Risk Mobilization", "rationale": "Traceable launch plan."}],
        "proposal_outline": {"sections": [{"sequence": 1, "proposal_section": "Executive Summary", "owner": "Proposal Lead"}]},
        "client_sections": [
            {"section_title": "Executive Summary", "body_markdown": "Draft executive summary.", "used_content_block_ids": [], "cited_requirement_codes": []},
            {"section_title": "Technical Approach", "body_markdown": "### Service Delivery Model\n- Maintain quality controls.", "used_content_block_ids": [], "cited_requirement_codes": []},
        ],
        "draft_package": {
            "sections": [
                {"section_title": "Executive Summary", "body_markdown": "Draft executive summary.", "used_content_block_ids": [], "cited_requirement_codes": []}
            ],
            "unresolved_items": ["Confirm pricing inputs."],
            "editor_notes": "Use the support bundle to finish pricing inputs.",
        },
        "compliance_matrix": [
            {
                "requirement_code": "REQ-001",
                "requirement_text": "The contractor shall provide a staffing plan.",
                "requirement_type": "COMPLIANCE_REQUIRED",
                "proposal_section": "Management and Staffing Plan",
                "owner": "Operations Lead",
                "status": "IN_PROGRESS",
            }
        ],
    }

    manifest = build_export_bundle(workspace)
    assert Path(manifest["client_docx_path"]).exists()
    assert Path(manifest["support_bundle_zip_path"]).exists()
    assert Path(manifest["docx_path"]).exists()
    assert Path(manifest["zip_path"]).exists()
    assert Path(manifest["markdown_path"]).exists()
    assert Path(manifest["html_preview_path"]).exists()
    assert Path(manifest["missing_input_checklist_path"]).exists()

    markdown = Path(manifest["markdown_path"]).read_text(encoding="utf-8")
    assert "Draft Proposal" in markdown
    assert "Proposal Starter" not in markdown
    html_preview = Path(manifest["html_preview_path"]).read_text(encoding="utf-8")
    assert "Proposal Response Draft" in html_preview
    checklist = Path(manifest["missing_input_checklist_path"]).read_text(encoding="utf-8")
    assert "Missing Input Checklist" in checklist


def test_template_validation_and_rendered_export_use_selected_template(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr(export_module, "EXPORTS_DIR", tmp_path)

    template_path = tmp_path / "client-template.docx"
    document = Document()
    document.add_paragraph("{{ opportunity_name }}")
    document.add_paragraph("{{ transmittal_letter }}")
    document.add_paragraph("{{ executive_summary }}")
    document.add_paragraph("{{ technical_approach }}")
    document.add_paragraph("{{ management_staffing_plan }}")
    document.add_paragraph("{{ transition_mobilization_plan }}")
    document.add_paragraph("{{ past_performance }}")
    document.add_paragraph("{{ pricing_and_commercials }}")
    document.add_paragraph("{{ compliance_attachments }}")
    document.save(template_path)

    report = validate_proposal_template(template_path)
    assert report["status"] == "VALID"
    assert not report["errors"]

    workspace = {
        "opportunity_id": "opp-demo",
        "opportunity_name": "Airport Terminal Janitorial and Day Porter Services",
        "client_name": "Metro Regional Airport Authority",
        "solicitation_id": "sol-demo",
        "extracted_deadline": "2026-04-10",
        "opportunity_summary": {
            "client_name": "Metro Regional Airport Authority",
            "opportunity_name": "Airport Terminal Janitorial and Day Porter Services",
            "solicitation_number": "MR-AA-2026-041",
            "proposal_due_date": "2026-04-10",
            "proposal_due_time": "2:00 PM ET",
        },
        "client_sections": [
            {"section_title": "Transmittal Letter", "body_markdown": "Transmittal text.", "used_content_block_ids": [], "cited_requirement_codes": []},
            {"section_title": "Executive Summary", "body_markdown": "Executive summary text.", "used_content_block_ids": [], "cited_requirement_codes": []},
            {"section_title": "Technical Approach", "body_markdown": "Technical approach text.", "used_content_block_ids": [], "cited_requirement_codes": []},
            {"section_title": "Management and Staffing Plan", "body_markdown": "Staffing plan text.", "used_content_block_ids": [], "cited_requirement_codes": []},
            {"section_title": "Transition and Mobilization Plan", "body_markdown": "Transition text.", "used_content_block_ids": [], "cited_requirement_codes": []},
            {"section_title": "Past Performance", "body_markdown": "Past performance text.", "used_content_block_ids": [], "cited_requirement_codes": []},
            {"section_title": "Pricing and Commercials", "body_markdown": "Pricing text.", "used_content_block_ids": [], "cited_requirement_codes": []},
            {"section_title": "Compliance and Attachments", "body_markdown": "Compliance text.", "used_content_block_ids": [], "cited_requirement_codes": []},
        ],
        "draft_package": {"sections": [], "unresolved_items": [], "editor_notes": ""},
        "compliance_matrix": [],
        "win_themes": [],
        "proposal_outline": {"sections": []},
    }

    manifest = build_export_bundle(
        workspace,
        proposal_template={
            "id": "template-1",
            "name": "Authority Proposal Template",
            "storage_path": template_path.as_posix(),
        },
    )
    assert manifest["proposal_template_id"] == "template-1"
    assert manifest["proposal_template_name"] == "Authority Proposal Template"

    rendered = Document(manifest["client_docx_path"])
    rendered_text = "\n".join(paragraph.text for paragraph in rendered.paragraphs)
    assert "Airport Terminal Janitorial and Day Porter Services" in rendered_text
    assert "Executive summary text." in rendered_text
    assert "Compliance text." in rendered_text


def test_template_render_replaces_footer_metadata_without_destroying_other_runs(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr(export_module, "EXPORTS_DIR", tmp_path)

    template_path = tmp_path / "footer-template.docx"
    document = Document()
    document.add_paragraph("{{ transmittal_letter }}")
    document.add_paragraph("{{ executive_summary }}")
    document.add_paragraph("{{ technical_approach }}")
    document.add_paragraph("{{ management_staffing_plan }}")
    document.add_paragraph("{{ transition_mobilization_plan }}")
    document.add_paragraph("{{ past_performance }}")
    document.add_paragraph("{{ pricing_and_commercials }}")
    document.add_paragraph("{{ compliance_attachments }}")
    footer = document.sections[0].footer
    footer_para = footer.paragraphs[0]
    footer_para.add_run("Prepared for ")
    footer_para.add_run("{{ client_name }}")
    footer_para.add_run(" | Proposal Due ")
    footer_para.add_run("{{ proposal_due }}")
    document.save(template_path)

    workspace = {
        "opportunity_id": "opp-footer-demo",
        "opportunity_name": "Enterprise Security Response Shell",
        "client_name": "St. Moritz Security Services, Inc.",
        "solicitation_id": "SMSSI-DEMO-2026-01",
        "opportunity_summary": {
            "client_name": "St. Moritz Security Services, Inc.",
            "opportunity_name": "Enterprise Security Response Shell",
            "solicitation_number": "SMSSI-DEMO-2026-01",
            "proposal_due_date": "April 2, 2026",
            "proposal_due_time": "5:00 PM ET",
        },
        "client_sections": [
            {"section_title": "Transmittal Letter", "body_markdown": "Transmittal text.", "used_content_block_ids": [], "cited_requirement_codes": []},
            {"section_title": "Executive Summary", "body_markdown": "Executive summary text.", "used_content_block_ids": [], "cited_requirement_codes": []},
            {"section_title": "Technical Approach", "body_markdown": "Technical approach text.", "used_content_block_ids": [], "cited_requirement_codes": []},
            {"section_title": "Management and Staffing Plan", "body_markdown": "Staffing plan text.", "used_content_block_ids": [], "cited_requirement_codes": []},
            {"section_title": "Transition and Mobilization Plan", "body_markdown": "Transition text.", "used_content_block_ids": [], "cited_requirement_codes": []},
            {"section_title": "Past Performance", "body_markdown": "Past performance text.", "used_content_block_ids": [], "cited_requirement_codes": []},
            {"section_title": "Pricing and Commercials", "body_markdown": "Pricing text.", "used_content_block_ids": [], "cited_requirement_codes": []},
            {"section_title": "Compliance and Attachments", "body_markdown": "Compliance text.", "used_content_block_ids": [], "cited_requirement_codes": []},
        ],
        "draft_package": {"sections": [], "unresolved_items": [], "editor_notes": ""},
        "compliance_matrix": [],
        "win_themes": [],
        "proposal_outline": {"sections": []},
    }

    manifest = build_export_bundle(
        workspace,
        proposal_template={
            "id": "template-footer",
            "name": "Footer Metadata Template",
            "storage_path": template_path.as_posix(),
        },
    )

    rendered = Document(manifest["client_docx_path"])
    footer_text = "\n".join(paragraph.text for paragraph in rendered.sections[0].footer.paragraphs)
    assert "Prepared for St. Moritz Security Services, Inc." in footer_text
    assert "Proposal Due April 2, 2026 5:00 PM ET" in footer_text


def test_draft_stage_coercion_accepts_list_of_sections() -> None:
    artifact = _coerce_draft_stage_artifact(
        [
            {
                "section_title": "Executive Summary",
                "body_markdown": "Summary text.",
                "used_content_block_ids": [],
                "cited_requirement_codes": ["REQ-001"],
            }
        ]
    )

    assert artifact.draft_package.client_sections[0].section_title == "Executive Summary"
    assert artifact.draft_package.editor_notes
    assert "Recovered draft content" in artifact.warnings[0]
