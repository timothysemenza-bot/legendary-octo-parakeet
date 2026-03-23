from io import BytesIO

from openpyxl import Workbook
from pptx import Presentation

from app.modules.rfp_parser.document_reader import extract_text_from_upload_batch


def _workbook_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Pricing"
    sheet["A1"] = "Line Item"
    sheet["B1"] = "Rate"
    sheet["A2"] = "Mowing"
    sheet["B2"] = 125.5

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _presentation_bytes() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Pre-Quote Conference"
    slide.placeholders[1].text = "Quotes must be submitted through NJSTART."

    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


def test_extract_text_from_upload_batch_merges_supported_files_in_order() -> None:
    result = extract_text_from_upload_batch(
        [
            (
                "scope.txt",
                (
                    "Proposal due date is 2026-07-15.\n"
                    "The contractor shall provide a staffing plan.\n"
                ).encode("utf-8"),
            ),
            (
                "pricing.md",
                (
                    "Evaluation criteria include pricing.\n"
                    "The offeror must submit a pricing narrative.\n"
                ).encode("utf-8"),
            ),
        ]
    )

    assert result.parsed_files == ["scope.txt", "pricing.md"]
    assert result.skipped_files == []
    assert result.warnings == []
    assert result.source_filename == "intake-batch-2-files.txt"
    assert [row.parse_status for row in result.documents] == ["PARSED", "PARSED"]
    assert result.documents[0].source_filename == "scope.txt"
    assert result.documents[0].upload_order == 1
    assert "staffing plan" in (result.documents[0].content_text or "")
    assert result.documents[0].source_payload is not None
    assert result.documents[1].upload_order == 2
    assert "pricing narrative" in (result.documents[1].content_text or "")
    assert result.documents[1].source_payload is not None
    assert result.combined_text.index("Source file: scope.txt") < result.combined_text.index(
        "Source file: pricing.md"
    )


def test_extract_text_from_upload_batch_parses_xlsx_and_pptx_files() -> None:
    result = extract_text_from_upload_batch(
        [
            ("price-sheet.xlsx", _workbook_bytes()),
            ("pre-quote-deck.pptx", _presentation_bytes()),
        ]
    )

    assert result.parsed_files == ["price-sheet.xlsx", "pre-quote-deck.pptx"]
    assert result.skipped_files == []
    assert "Sheet: Pricing" in (result.documents[0].content_text or "")
    assert "A2=Mowing" in (result.documents[0].content_text or "")
    assert "Slide 1" in (result.documents[1].content_text or "")
    assert "Pre-Quote Conference" in (result.documents[1].content_text or "")
    assert result.combined_text.index("Source file: price-sheet.xlsx") < result.combined_text.index(
        "Source file: pre-quote-deck.pptx"
    )


def test_extract_text_from_upload_batch_skips_invalid_files_with_warnings() -> None:
    result = extract_text_from_upload_batch(
        [
            ("notes.txt", b"too short"),
            ("archive.zip", b"binary"),
            (
                "rfp.txt",
                (
                    "Evaluation criteria include technical approach.\n"
                    "The offeror shall provide a transition plan.\n"
                ).encode("utf-8"),
            ),
        ]
    )

    assert result.parsed_files == ["rfp.txt"]
    assert result.skipped_files == ["notes.txt", "archive.zip"]
    assert result.source_filename == "intake-batch-1-files.txt"
    assert [row.parse_status for row in result.documents] == ["SKIPPED", "SKIPPED", "PARSED"]
    assert result.documents[0].skip_reason == "extracted content is too short to parse."
    assert result.documents[0].content_text is None
    assert result.documents[0].source_payload == b"too short"
    assert result.documents[1].skip_reason == "Unsupported file type. Use TXT, MD, PDF, DOCX, XLSX, or PPTX."
    assert result.documents[1].content_text is None
    assert result.documents[1].source_payload == b"binary"
    assert result.documents[2].extracted_text_length > 20
    assert "transition plan" in (result.documents[2].content_text or "")
    assert result.documents[2].source_payload is not None
    assert any("too short to parse" in warning for warning in result.warnings)
    assert any("Unsupported file type" in warning for warning in result.warnings)
