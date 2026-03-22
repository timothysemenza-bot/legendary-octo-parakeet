from __future__ import annotations

import csv
import html
import json
import zipfile
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from docx import Document

from app.core.config import EXPORTS_DIR
from app.modules.proposal_builder.template_renderer import render_proposal_template


def _safe_slug(value: str) -> str:
    normalized = "".join(char.lower() if char.isalnum() else "-" for char in value).strip("-")
    return "-".join(part for part in normalized.split("-") if part)[:80] or "proposal-builder"


def _short_requirement_text(value: str, *, limit: int = 110) -> str:
    cleaned = " ".join((value or "").split())
    if len(cleaned) <= limit:
        return cleaned
    return cleaned[: limit - 3].rstrip(" ,;:.") + "..."


def _summary(workspace: dict[str, Any]) -> dict[str, Any]:
    return workspace.get("opportunity_summary") or {}


def _client_sections(workspace: dict[str, Any]) -> list[dict[str, Any]]:
    sections = workspace.get("client_sections") or []
    if sections:
        return sections
    draft_package = workspace.get("draft_package") or {}
    return draft_package.get("client_sections") or draft_package.get("sections") or []


def _title_parts(workspace: dict[str, Any]) -> tuple[str, str, str]:
    summary = _summary(workspace)
    opportunity_name = summary.get("opportunity_name") or workspace["opportunity_name"]
    client_name = summary.get("client_name") or workspace["client_name"]
    solicitation_number = summary.get("solicitation_number") or workspace.get("solicitation_id") or "Working Draft"
    return opportunity_name, client_name, solicitation_number


def _proposal_due_label(workspace: dict[str, Any]) -> str:
    summary = _summary(workspace)
    due_date = summary.get("proposal_due_date") or workspace.get("extracted_deadline") or "Pending confirmation"
    due_time = summary.get("proposal_due_time") or ""
    return f"{due_date} {due_time}".strip()


def _provider_name(workspace: dict[str, Any]) -> str:
    proposal_config = workspace.get("proposal_config") or {}
    return proposal_config.get("client_profile_name") or "Boss Key"


def _write_markdown(markdown_path: Path, content: str) -> None:
    markdown_path.write_text(content, encoding="utf-8")


def _write_json(json_path: Path, payload: dict[str, Any]) -> None:
    json_path.write_text(json.dumps(payload, indent=2, default=str), encoding="utf-8")


def _write_compliance_csv(csv_path: Path, rows: list[dict[str, Any]]) -> None:
    fieldnames = [
        "requirement_code",
        "requirement_text",
        "requirement_type",
        "proposal_section",
        "owner",
        "status",
    ]
    with csv_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows:
            writer.writerow({field: row.get(field, "") for field in fieldnames})


def _render_markdownish_html(body: str) -> str:
    parts: list[str] = []
    paragraph_lines: list[str] = []
    bullet_items: list[str] = []

    def flush_paragraph() -> None:
        if not paragraph_lines:
            return
        parts.append(f"<p>{html.escape(' '.join(paragraph_lines).strip())}</p>")
        paragraph_lines.clear()

    def flush_bullets() -> None:
        if not bullet_items:
            return
        parts.append("<ul>" + "".join(f"<li>{html.escape(item)}</li>" for item in bullet_items) + "</ul>")
        bullet_items.clear()

    for raw_line in body.splitlines():
        stripped = raw_line.strip()
        if not stripped:
            flush_paragraph()
            flush_bullets()
            continue
        if stripped.startswith("### "):
            flush_paragraph()
            flush_bullets()
            parts.append(f"<h3>{html.escape(stripped[4:].strip())}</h3>")
            continue
        if stripped.startswith("- "):
            flush_paragraph()
            bullet_items.append(stripped[2:].strip())
            continue
        flush_bullets()
        paragraph_lines.append(stripped)

    flush_paragraph()
    flush_bullets()
    return "\n".join(parts)


def _render_html_preview(workspace: dict[str, Any]) -> str:
    opportunity_name, client_name, solicitation_number = _title_parts(workspace)
    proposal_due = _proposal_due_label(workspace)
    provider_name = _provider_name(workspace)
    section_markup = []
    for section in _client_sections(workspace):
        section_markup.append(
            f"""
            <section class="section-card">
              <h2>{html.escape(section['section_title'])}</h2>
              {_render_markdownish_html(section['body_markdown'])}
            </section>
            """.strip()
        )
    return f"""<!doctype html>
<html lang="en">
  <head>
    <meta charset="utf-8">
    <title>{html.escape(opportunity_name)} - Proposal Preview</title>
    <style>
      :root {{
        --st-red: #D0001C;
        --st-deep-red: #A20502;
        --st-navy: #1A2536;
        --st-gold: #B57420;
        --st-gray: #E7E9EC;
      }}
      body {{
        margin: 0;
        background: #f4f5f7;
        color: var(--st-navy);
        font-family: Calibri, Arial, sans-serif;
      }}
      .shell {{
        max-width: 960px;
        margin: 0 auto;
        background: white;
        min-height: 100vh;
        box-shadow: 0 18px 60px rgba(26, 37, 54, 0.12);
      }}
      .top-band {{ height: 18px; background: var(--st-red); }}
      .hero {{
        padding: 36px 52px 28px;
        border-bottom: 2px solid var(--st-gray);
      }}
      .brand {{
        font-family: 'Arial Narrow', Arial, sans-serif;
        font-size: 14px;
        letter-spacing: 0.08em;
        text-transform: uppercase;
        color: var(--st-deep-red);
        margin-bottom: 8px;
      }}
      h1 {{
        font-family: 'Arial Narrow', Arial, sans-serif;
        font-size: 34px;
        line-height: 1.05;
        margin: 0 0 10px;
        color: var(--st-navy);
      }}
      .subtitle {{
        font-size: 18px;
        color: var(--st-deep-red);
        margin-bottom: 22px;
      }}
      .meta-grid {{
        display: grid;
        grid-template-columns: repeat(2, minmax(0, 1fr));
        gap: 16px 28px;
      }}
      .meta-label {{
        font-size: 11px;
        letter-spacing: 0.1em;
        text-transform: uppercase;
        color: #5b6472;
        margin-bottom: 4px;
      }}
      .meta-value {{
        font-size: 16px;
        color: var(--st-navy);
      }}
      .content {{
        padding: 28px 52px 40px;
      }}
      .section-card {{
        margin-bottom: 22px;
        padding-bottom: 18px;
        border-bottom: 1px solid var(--st-gray);
      }}
      .section-card h2 {{
        font-family: 'Arial Narrow', Arial, sans-serif;
        font-size: 24px;
        margin: 0 0 12px;
        color: var(--st-navy);
      }}
      .section-card h3 {{
        font-size: 16px;
        color: var(--st-deep-red);
        margin: 16px 0 8px;
      }}
      .section-card p, .section-card li {{
        font-size: 15px;
        line-height: 1.5;
      }}
      .section-card ul {{
        margin: 8px 0 12px 20px;
      }}
      .gaps {{
        margin: 8px 0 28px;
        padding: 16px 18px;
        border-left: 4px solid var(--st-gold);
        background: #fbf7ef;
      }}
      .gaps h2 {{
        margin-top: 0;
      }}
      footer {{
        padding: 14px 52px 22px;
        border-top: 3px solid var(--st-navy);
        font-size: 12px;
        color: #5b6472;
      }}
    </style>
  </head>
  <body>
    <main class="shell">
      <div class="top-band"></div>
      <section class="hero">
        <div class="brand">{html.escape(provider_name)}</div>
        <h1>{html.escape(opportunity_name)}</h1>
        <div class="subtitle">Proposal Response Draft</div>
        <div class="meta-grid">
          <div><div class="meta-label">Prepared For</div><div class="meta-value">{html.escape(client_name)}</div></div>
          <div><div class="meta-label">Solicitation</div><div class="meta-value">{html.escape(solicitation_number)}</div></div>
          <div><div class="meta-label">Proposal Due</div><div class="meta-value">{html.escape(proposal_due)}</div></div>
          <div><div class="meta-label">Generated</div><div class="meta-value">{html.escape(datetime.now(UTC).strftime('%B %d, %Y %I:%M %p UTC'))}</div></div>
        </div>
      </section>
      <section class="content">
        <section class="gaps">
          <h2>Known Gaps</h2>
          <ul>
            {''.join(f"<li>{html.escape(item)}</li>" for item in (workspace.get('document_gaps') or []))}
          </ul>
        </section>
        {''.join(section_markup)}
      </section>
      <footer>{html.escape(provider_name)} | Prepared for {html.escape(client_name)}</footer>
    </main>
  </body>
</html>
"""


def _write_missing_input_checklist(path: Path, workspace: dict[str, Any]) -> None:
    items = []
    for group_name in ("document_gaps", "setup_gaps", "warnings", "assumption_flags"):
        for item in workspace.get(group_name, []) or []:
            cleaned = str(item).strip()
            if cleaned and cleaned not in items:
                items.append(cleaned)
    lines = ["# Missing Input Checklist", ""]
    if not items:
        lines.append("- No open input gaps were recorded at export time.")
    else:
        for item in items:
            lines.append(f"- {item}")
    lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")


def _render_markdown_bundle(workspace: dict[str, Any]) -> str:
    opportunity_name, client_name, solicitation_number = _title_parts(workspace)
    lines = [
        f"# {opportunity_name} Draft Proposal",
        "",
        f"- Buyer: {client_name}",
        f"- Solicitation: {solicitation_number}",
        f"- Proposal Due: {_proposal_due_label(workspace)}",
        "",
    ]

    for section in _client_sections(workspace):
        lines.append(f"## {section['section_title']}")
        lines.append("")
        lines.append(section["body_markdown"].strip())
        lines.append("")

    lines.append("## Compliance Appendix")
    lines.append("")
    for row in workspace.get("compliance_matrix", []):
        lines.append(
            f"- **{row.get('requirement_code', '')}** ({row.get('proposal_section', '')}) {row.get('requirement_text', '')}"
        )
    lines.append("")
    return "\n".join(lines).strip() + "\n"


def _add_markdownish_body(document: Document, body: str) -> None:
    paragraph_lines: list[str] = []

    def flush_paragraph() -> None:
        if not paragraph_lines:
            return
        document.add_paragraph(" ".join(paragraph_lines).strip())
        paragraph_lines.clear()

    for raw_line in body.splitlines():
        line = raw_line.rstrip()
        stripped = line.strip()
        if not stripped:
            flush_paragraph()
            continue
        if stripped.startswith("### "):
            flush_paragraph()
            document.add_heading(stripped[4:].strip(), level=2)
            continue
        if stripped.startswith("- "):
            flush_paragraph()
            document.add_paragraph(stripped[2:].strip(), style="List Bullet")
            continue
        paragraph_lines.append(stripped)

    flush_paragraph()


def _build_docx(workspace: dict[str, Any], target_path: Path) -> None:
    opportunity_name, client_name, solicitation_number = _title_parts(workspace)
    document = Document()
    document.add_heading(f"{opportunity_name} Draft Proposal", level=0)
    document.add_paragraph(f"Buyer: {client_name}")
    document.add_paragraph(f"Solicitation: {solicitation_number}")
    document.add_paragraph(f"Proposal Due: {_proposal_due_label(workspace)}")

    for section in _client_sections(workspace):
        document.add_heading(section["section_title"], level=1)
        _add_markdownish_body(document, section["body_markdown"])

    document.add_heading("Compliance Appendix", level=1)
    table = document.add_table(rows=1, cols=6)
    headers = table.rows[0].cells
    headers[0].text = "Requirement"
    headers[1].text = "Requirement Summary"
    headers[2].text = "Type"
    headers[3].text = "Proposal Section"
    headers[4].text = "Owner"
    headers[5].text = "Status"
    for row in workspace.get("compliance_matrix", []):
        cells = table.add_row().cells
        cells[0].text = row.get("requirement_code", "")
        cells[1].text = _short_requirement_text(row.get("requirement_text", ""))
        cells[2].text = row.get("requirement_type", "")
        cells[3].text = row.get("proposal_section", "")
        cells[4].text = row.get("owner", "")
        cells[5].text = row.get("status", "")

    document.save(target_path)


def build_export_bundle(workspace: dict[str, Any], proposal_template: dict[str, Any] | None = None) -> dict[str, Any]:
    export_root = EXPORTS_DIR / workspace["opportunity_id"]
    export_root.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now(UTC).strftime("%Y%m%d-%H%M%S")
    export_dir = export_root / f"{timestamp}-{_safe_slug(workspace['opportunity_name'])}"
    export_dir.mkdir(parents=True, exist_ok=True)

    markdown_path = export_dir / "proposal-draft.md"
    json_path = export_dir / "proposal-builder-workspace.json"
    csv_path = export_dir / "compliance-matrix.csv"
    html_preview_path = export_dir / "proposal-preview.html"
    missing_input_checklist_path = export_dir / "missing-input-checklist.md"
    client_docx_path = export_dir / "client-proposal-draft.docx"
    bundle_zip_path = export_dir / "proposal-builder-support-bundle.zip"

    markdown = _render_markdown_bundle(workspace)
    _write_markdown(markdown_path, markdown)
    _write_json(json_path, workspace)
    _write_compliance_csv(csv_path, workspace.get("compliance_matrix", []))
    html_preview_path.write_text(_render_html_preview(workspace), encoding="utf-8")
    _write_missing_input_checklist(missing_input_checklist_path, workspace)
    if proposal_template and proposal_template.get("storage_path"):
        render_proposal_template(
            workspace=workspace,
            template_path=Path(str(proposal_template["storage_path"])),
            output_path=client_docx_path,
        )
    else:
        _build_docx(workspace, client_docx_path)

    with zipfile.ZipFile(bundle_zip_path, "w", compression=zipfile.ZIP_DEFLATED) as bundle:
        for path in (markdown_path, json_path, csv_path, html_preview_path, missing_input_checklist_path, client_docx_path):
            bundle.write(path, arcname=path.name)

    return {
        "export_directory": export_dir.as_posix(),
        "client_docx_path": client_docx_path.as_posix(),
        "support_bundle_zip_path": bundle_zip_path.as_posix(),
        "proposal_template_id": proposal_template.get("id") if proposal_template else None,
        "proposal_template_name": proposal_template.get("name") if proposal_template else None,
        "docx_path": client_docx_path.as_posix(),
        "markdown_path": markdown_path.as_posix(),
        "json_path": json_path.as_posix(),
        "compliance_csv_path": csv_path.as_posix(),
        "html_preview_path": html_preview_path.as_posix(),
        "missing_input_checklist_path": missing_input_checklist_path.as_posix(),
        "zip_path": bundle_zip_path.as_posix(),
        "generated_at": datetime.now(UTC),
    }


def _write_review_report(path: Path, findings: list[dict[str, Any]]) -> None:
    lines = ["# Review Findings", ""]
    for finding in findings:
        lines.append(f"## {finding.get('review_stage', 'review')}: {finding.get('title', 'Finding')}")
        lines.append(f"- Severity: {finding.get('severity', '')}")
        lines.append(f"- Disposition: {finding.get('disposition', '')}")
        lines.append(f"- Recommendation: {finding.get('recommendation', '')}")
        if finding.get("section_title"):
            lines.append(f"- Section: {finding.get('section_title', '')}")
        lines.append("")
    path.write_text("\n".join(lines).strip() + "\n", encoding="utf-8")


def _write_production_checklist(path: Path, package_run: dict[str, Any]) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["Stage", "Status", "Approval", "Blocking Issues"])
        for stage in package_run.get("stages", []):
            writer.writerow(
                [
                    stage.get("stage_name", ""),
                    stage.get("status", ""),
                    stage.get("approval_status", "") or stage.get("approval_required", "") or "",
                    "; ".join(stage.get("blocking_issues", []) or []),
                ]
            )


def _write_attachments_manifest(path: Path, package_run: dict[str, Any]) -> list[str]:
    form_package = package_run.get("form_package") or {}
    files = [
        *(form_package.get("resumes_packet_paths") or []),
        *(form_package.get("references_packet_paths") or []),
    ]
    path.write_text(json.dumps({"files": files}, indent=2), encoding="utf-8")
    return files


def build_proposal_package_export_bundle(
    *,
    workspace: dict[str, Any],
    package_run: dict[str, Any],
    target_dir: Path,
    proposal_template: dict[str, Any] | None = None,
) -> dict[str, Any]:
    target_dir.mkdir(parents=True, exist_ok=True)

    markdown_path = target_dir / "proposal-package-draft.md"
    json_path = target_dir / "proposal-package-workspace.json"
    csv_path = target_dir / "compliance-matrix.csv"
    client_docx_path = target_dir / "final-proposal-package.docx"
    review_report_path = target_dir / "review-report.md"
    production_checklist_path = target_dir / "production-checklist.csv"
    attachments_manifest_path = target_dir / "attachments-manifest.json"
    support_bundle_zip_path = target_dir / "proposal-package-support-bundle.zip"
    attachments_bundle_path = target_dir / "attachments-bundle.zip"

    markdown = _render_markdown_bundle(workspace)
    _write_markdown(markdown_path, markdown)
    _write_json(json_path, {"workspace": workspace, "package_run": package_run})
    _write_compliance_csv(csv_path, workspace.get("compliance_matrix", []))
    if proposal_template and proposal_template.get("storage_path"):
        render_proposal_template(
            workspace=workspace,
            template_path=Path(str(proposal_template["storage_path"])),
            output_path=client_docx_path,
        )
    else:
        _build_docx(workspace, client_docx_path)
    _write_review_report(review_report_path, package_run.get("review_findings", []))
    _write_production_checklist(production_checklist_path, package_run)
    attachment_files = _write_attachments_manifest(attachments_manifest_path, package_run)

    pricing_workbook_path = ((package_run.get("pricing_package") or {}).get("workbook_path") or "").strip()
    forms_bundle_path = ""
    for item in (package_run.get("form_package") or {}).get("completed_forms", []):
        if item.get("label") == "Completed Forms Bundle":
            forms_bundle_path = str(item.get("path", "")).strip()
            break

    with zipfile.ZipFile(attachments_bundle_path, "w", compression=zipfile.ZIP_DEFLATED) as bundle:
        for raw_path in attachment_files:
            file_path = Path(raw_path)
            if file_path.exists():
                bundle.write(file_path, arcname=file_path.name)

    with zipfile.ZipFile(support_bundle_zip_path, "w", compression=zipfile.ZIP_DEFLATED) as bundle:
        for path in (
            markdown_path,
            json_path,
            csv_path,
            client_docx_path,
            review_report_path,
            production_checklist_path,
            attachments_manifest_path,
            attachments_bundle_path,
        ):
            bundle.write(path, arcname=path.name)
        for raw_path in (pricing_workbook_path, forms_bundle_path):
            if raw_path:
                file_path = Path(raw_path)
                if file_path.exists():
                    bundle.write(file_path, arcname=file_path.name)

    return {
        "export_directory": target_dir.as_posix(),
        "client_docx_path": client_docx_path.as_posix(),
        "support_bundle_zip_path": support_bundle_zip_path.as_posix(),
        "proposal_template_id": proposal_template.get("id") if proposal_template else None,
        "proposal_template_name": proposal_template.get("name") if proposal_template else None,
        "docx_path": client_docx_path.as_posix(),
        "markdown_path": markdown_path.as_posix(),
        "json_path": json_path.as_posix(),
        "compliance_csv_path": csv_path.as_posix(),
        "zip_path": support_bundle_zip_path.as_posix(),
        "pricing_workbook_path": pricing_workbook_path or None,
        "forms_bundle_path": forms_bundle_path or None,
        "attachments_bundle_path": attachments_bundle_path.as_posix(),
        "review_report_path": review_report_path.as_posix(),
        "production_checklist_path": production_checklist_path.as_posix(),
        "generated_at": datetime.now(UTC),
    }
