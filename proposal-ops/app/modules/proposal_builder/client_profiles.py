from __future__ import annotations

import csv
import io
import json
import re
import uuid
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core.audit import log_audit_event
from app.core.config import ARTIFACTS_DIR
from app.modules.opportunity_intake.models import Opportunity
from app.modules.proposal_builder.models import (
    ClientOnboardingPack,
    ClientPricingModel,
    ClientProfile,
    ClientPlaybook,
    ClientProposalTemplate,
    OpportunityProposalConfig,
)
from app.modules.proposal_builder.schemas import (
    ClientPricingModelCanonicalArtifact,
    ClientPricingModelCreateRequest,
    ClientPricingModelOption,
    ClientPricingModelResponse,
    ClientProfileCreateRequest,
    ClientProfileDefaultsRequest,
    ClientProfileOption,
    ClientProfileResponse,
    ClientProposalTemplateOption,
    ClientProposalTemplateResponse,
    OpportunityProposalConfigArtifact,
    OpportunityProposalConfigSelectionRequest,
    PricingModelLaborCategoryArtifact,
    ValidationReportArtifact,
)
from app.modules.proposal_builder.template_renderer import REQUIRED_PROPOSAL_TEMPLATE_ANCHORS, validate_proposal_template


CLIENT_PROFILE_STORAGE_DIR = ARTIFACTS_DIR / "client_profiles"
CLIENT_PROFILE_STORAGE_DIR.mkdir(parents=True, exist_ok=True)
HEADER_ALIASES = {
    "code": {"code", "labor category", "labor_category", "category", "role", "labor code"},
    "label": {"label", "title", "name", "labor label", "labor_title"},
    "hourly_rate": {"hourly_rate", "hourly rate", "pay_rate", "pay rate", "rate"},
    "burden_factor": {"burden_factor", "burden factor", "burden"},
    "markup_factor": {"markup_factor", "markup factor", "markup", "multiplier"},
    "default_hours_per_week": {
        "default_hours_per_week",
        "default hours per week",
        "hours_per_week",
        "hours/week",
        "weekly_hours",
    },
}
DEFAULT_CATEGORY_HOURS = {
    "janitor": 160.0,
    "day_porter": 80.0,
    "supervisor": 20.0,
    "project_manager": 6.0,
}


class ClientProfileError(RuntimeError):
    pass


@dataclass
class ResolvedOpportunityConfig:
    config_row: OpportunityProposalConfig
    config: OpportunityProposalConfigArtifact
    available_profiles: list[ClientProfileOption]
    available_pricing_models: list[ClientPricingModelOption]
    available_proposal_templates: list[ClientProposalTemplateOption]


def _parse_json(raw: str | None, default: Any) -> Any:
    if not raw:
        return default
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        return default


def _dump_json(value: Any) -> str:
    return json.dumps(value, default=str)


def _normalize_space(value: str | None) -> str:
    return re.sub(r"\s+", " ", (value or "").strip())


def _normalize_key(value: str | None) -> str:
    cleaned = _normalize_space(value).lower()
    cleaned = re.sub(r"[^a-z0-9]+", "_", cleaned).strip("_")
    return cleaned


def _unique_strings(values: list[str] | tuple[str, ...]) -> list[str]:
    ordered: list[str] = []
    for value in values:
        cleaned = _normalize_space(value)
        if cleaned and cleaned not in ordered:
            ordered.append(cleaned)
    return ordered


def _safe_slug(value: str) -> str:
    normalized = "".join(char.lower() if char.isalnum() else "-" for char in value).strip("-")
    return "-".join(part for part in normalized.split("-") if part)[:80] or "client-profile"


def _header_key(raw: str) -> str:
    candidate = _normalize_key(raw)
    for canonical, aliases in HEADER_ALIASES.items():
        if candidate in {_normalize_key(item) for item in aliases}:
            return canonical
    return candidate


def _default_hours_for_code(code: str) -> float:
    lowered = _normalize_key(code)
    if lowered in DEFAULT_CATEGORY_HOURS:
        return DEFAULT_CATEGORY_HOURS[lowered]
    if "porter" in lowered:
        return DEFAULT_CATEGORY_HOURS["day_porter"]
    if "manager" in lowered:
        return DEFAULT_CATEGORY_HOURS["project_manager"]
    if "supervisor" in lowered or "lead" in lowered:
        return DEFAULT_CATEGORY_HOURS["supervisor"]
    if "janitor" in lowered or "clean" in lowered:
        return DEFAULT_CATEGORY_HOURS["janitor"]
    return 0.0


def _storage_dir_for_profile(profile_id: str, *parts: str) -> Path:
    target = CLIENT_PROFILE_STORAGE_DIR / profile_id
    for part in parts:
        target = target / part
    target.mkdir(parents=True, exist_ok=True)
    return target


def _validation_report(payload: dict[str, Any] | None) -> ValidationReportArtifact | None:
    if not payload:
        return None
    return ValidationReportArtifact.model_validate(payload)


def _serialize_pricing_model(model: ClientPricingModel) -> ClientPricingModelResponse:
    return ClientPricingModelResponse(
        id=model.id,
        client_profile_id=model.client_profile_id,
        name=model.name,
        status=model.status,
        onboarding_mode=model.onboarding_mode,
        source_format=model.source_format,
        source_filename=model.source_filename,
        source_file_path=model.source_file_path,
        workbook_template_filename=model.workbook_template_filename,
        workbook_template_path=model.workbook_template_path,
        canonical_model=ClientPricingModelCanonicalArtifact.model_validate(_parse_json(model.canonical_model_json, {}))
        if model.canonical_model_json
        else None,
        validation_report=_validation_report(_parse_json(model.validation_report_json, {})),
        created_at=model.created_at,
        updated_at=model.updated_at,
    )


def _serialize_template(template: ClientProposalTemplate) -> ClientProposalTemplateResponse:
    return ClientProposalTemplateResponse(
        id=template.id,
        client_profile_id=template.client_profile_id,
        name=template.name,
        status=template.status,
        source_filename=template.source_filename,
        storage_path=template.storage_path,
        anchor_map=_parse_json(template.anchor_map_json, {}),
        validation_report=_validation_report(_parse_json(template.validation_report_json, {})),
        created_at=template.created_at,
        updated_at=template.updated_at,
    )


class ClientProfileService:
    def __init__(self, db: Session) -> None:
        self.db = db

    def _profile(self, profile_id: str) -> ClientProfile:
        profile = self.db.get(ClientProfile, profile_id)
        if profile is None:
            raise ClientProfileError("Client profile not found.")
        return profile

    def _pricing_model(self, pricing_model_id: str) -> ClientPricingModel:
        model = self.db.get(ClientPricingModel, pricing_model_id)
        if model is None:
            raise ClientProfileError("Pricing model not found.")
        return model

    def _proposal_template(self, proposal_template_id: str) -> ClientProposalTemplate:
        template = self.db.get(ClientProposalTemplate, proposal_template_id)
        if template is None:
            raise ClientProfileError("Proposal template not found.")
        return template

    def _opportunity(self, opportunity_id: str) -> Opportunity:
        opportunity = self.db.get(Opportunity, opportunity_id)
        if opportunity is None:
            raise ClientProfileError("Opportunity not found.")
        return opportunity

    def _config_row(self, opportunity_id: str) -> OpportunityProposalConfig | None:
        stmt = select(OpportunityProposalConfig).where(OpportunityProposalConfig.opportunity_id == opportunity_id)
        return self.db.scalars(stmt).first()

    def _profile_aliases(self, profile: ClientProfile) -> list[str]:
        return _parse_json(profile.aliases_json, [])

    def _profile_matches_name(self, profile: ClientProfile, client_name: str) -> bool:
        normalized_target = _normalize_key(client_name)
        if not normalized_target:
            return False
        candidates = [profile.display_name, *self._profile_aliases(profile)]
        return normalized_target in {_normalize_key(value) for value in candidates}

    def _list_profiles(self) -> list[ClientProfile]:
        stmt = select(ClientProfile).order_by(func.lower(ClientProfile.display_name).asc())
        return list(self.db.scalars(stmt))

    def _list_pricing_models(self, profile_id: str) -> list[ClientPricingModel]:
        stmt = (
            select(ClientPricingModel)
            .where(ClientPricingModel.client_profile_id == profile_id)
            .order_by(ClientPricingModel.updated_at.desc(), ClientPricingModel.created_at.desc())
        )
        return list(self.db.scalars(stmt))

    def _list_proposal_templates(self, profile_id: str) -> list[ClientProposalTemplate]:
        stmt = (
            select(ClientProposalTemplate)
            .where(ClientProposalTemplate.client_profile_id == profile_id)
            .order_by(ClientProposalTemplate.updated_at.desc(), ClientProposalTemplate.created_at.desc())
        )
        return list(self.db.scalars(stmt))

    def _serialize_profile(self, profile: ClientProfile, *, include_assets: bool = True) -> ClientProfileResponse:
        pricing_models: list[ClientPricingModelResponse] = []
        proposal_templates: list[ClientProposalTemplateResponse] = []
        default_pricing_model_name = None
        default_proposal_template_name = None
        active_playbook_version = None
        if include_assets:
            pricing_models = [_serialize_pricing_model(model) for model in self._list_pricing_models(profile.id)]
            proposal_templates = [_serialize_template(template) for template in self._list_proposal_templates(profile.id)]
        if profile.default_pricing_model_id:
            model = self.db.get(ClientPricingModel, profile.default_pricing_model_id)
            default_pricing_model_name = model.name if model else None
        if profile.default_proposal_template_id:
            template = self.db.get(ClientProposalTemplate, profile.default_proposal_template_id)
            default_proposal_template_name = template.name if template else None
        if profile.active_playbook_id:
            playbook = self.db.get(ClientPlaybook, profile.active_playbook_id)
            active_playbook_version = playbook.version if playbook else None
        return ClientProfileResponse(
            id=profile.id,
            display_name=profile.display_name,
            aliases=self._profile_aliases(profile),
            approved_content_tags=_parse_json(profile.approved_content_tags_json, []),
            region_tags=_parse_json(profile.region_tags_json, []),
            service_tags=_parse_json(profile.service_tags_json, []),
            status=profile.status,
            active_onboarding_pack_id=profile.active_onboarding_pack_id,
            active_playbook_version=active_playbook_version,
            client_environment_status=profile.client_environment_status,
            default_pricing_model_id=profile.default_pricing_model_id,
            default_proposal_template_id=profile.default_proposal_template_id,
            default_pricing_model_name=default_pricing_model_name,
            default_proposal_template_name=default_proposal_template_name,
            pricing_models=pricing_models,
            proposal_templates=proposal_templates,
            created_at=profile.created_at,
            updated_at=profile.updated_at,
        )

    def _profile_options(self) -> list[ClientProfileOption]:
        return [
            ClientProfileOption(id=profile.id, display_name=profile.display_name, status=profile.status)
            for profile in self._list_profiles()
            if profile.status.upper() == "ACTIVE"
        ]

    def _pricing_model_options(self, profile_id: str | None) -> list[ClientPricingModelOption]:
        if not profile_id:
            return []
        return [
            ClientPricingModelOption(
                id=model.id,
                name=model.name,
                status=model.status,
                onboarding_mode=model.onboarding_mode,
                source_format=model.source_format,
            )
            for model in self._list_pricing_models(profile_id)
        ]

    def _proposal_template_options(self, profile_id: str | None) -> list[ClientProposalTemplateOption]:
        if not profile_id:
            return []
        options: list[ClientProposalTemplateOption] = []
        for template in self._list_proposal_templates(profile_id):
            report = _validation_report(_parse_json(template.validation_report_json, {}))
            options.append(
                ClientProposalTemplateOption(
                    id=template.id,
                    name=template.name,
                    status=template.status,
                    source_filename=template.source_filename,
                    valid=bool(report and not report.errors and template.status.upper() == "APPROVED"),
                )
            )
        return options

    def list_profiles(self) -> list[ClientProfileResponse]:
        return [self._serialize_profile(profile, include_assets=False) for profile in self._list_profiles()]

    def get_profile(self, profile_id: str) -> ClientProfileResponse:
        return self._serialize_profile(self._profile(profile_id), include_assets=True)

    def create_profile(self, payload: ClientProfileCreateRequest) -> ClientProfileResponse:
        profile = ClientProfile(
            display_name=_normalize_space(payload.display_name),
            aliases_json=_dump_json(_unique_strings(payload.aliases)),
            approved_content_tags_json=_dump_json(_unique_strings(payload.approved_content_tags)),
            region_tags_json=_dump_json(_unique_strings(payload.region_tags)),
            service_tags_json=_dump_json(_unique_strings(payload.service_tags)),
            status=payload.status.upper(),
        )
        self.db.add(profile)
        self.db.flush()
        log_audit_event(
            self.db,
            opportunity_id=None,
            actor="admin",
            action="client_profile_created",
            after_state_json=_dump_json({"client_profile_id": profile.id, "display_name": profile.display_name}),
        )
        self.db.commit()
        self.db.refresh(profile)
        return self._serialize_profile(profile, include_assets=True)

    def _canonical_from_json(self, payload: dict[str, Any]) -> tuple[ClientPricingModelCanonicalArtifact | None, ValidationReportArtifact]:
        errors: list[str] = []
        warnings: list[str] = []
        source = payload
        if "canonical_model" in source and isinstance(source["canonical_model"], dict):
            source = source["canonical_model"]
        rows = source.get("labor_categories") if isinstance(source.get("labor_categories"), list) else source.get("line_items", [])
        normalized_rows: list[PricingModelLaborCategoryArtifact] = []
        for item in rows if isinstance(rows, list) else []:
            if not isinstance(item, dict):
                continue
            code = _normalize_key(str(item.get("code") or item.get("labor_category") or item.get("label") or item.get("name") or ""))
            if not code:
                continue
            label = _normalize_space(str(item.get("label") or item.get("name") or code.replace("_", " ").title()))
            try:
                hourly_rate = float(item.get("hourly_rate", item.get("rate", 0)))
                burden_factor = float(item.get("burden_factor", item.get("burden", 0)))
                markup_factor = float(item.get("markup_factor", item.get("markup", 0)))
                default_hours = float(item.get("default_hours_per_week", item.get("hours_per_week", _default_hours_for_code(code))))
            except (TypeError, ValueError):
                errors.append(f"Pricing row '{label}' has invalid numeric values.")
                continue
            if default_hours == 0:
                warnings.append(f"Pricing row '{label}' does not include default hours; manual review is required.")
            normalized_rows.append(
                PricingModelLaborCategoryArtifact(
                    code=code,
                    label=label,
                    hourly_rate=hourly_rate,
                    burden_factor=burden_factor,
                    markup_factor=markup_factor,
                    default_hours_per_week=default_hours,
                    apply_site_multiplier=bool(item.get("apply_site_multiplier", True)),
                    apply_continuous_coverage_multiplier=bool(item.get("apply_continuous_coverage_multiplier", True)),
                )
            )
        if not normalized_rows:
            errors.append("No labor categories were found in the JSON payload.")
        canonical = None
        if not errors:
            canonical = ClientPricingModelCanonicalArtifact(
                labor_categories=normalized_rows,
                staffing_rules={
                    "site_multiplier": float(source.get("staffing_rules", {}).get("site_multiplier", source.get("site_multiplier", 1.0))),
                    "continuous_coverage_multiplier": float(
                        source.get("staffing_rules", {}).get("continuous_coverage_multiplier", source.get("continuous_coverage_multiplier", 1.35))
                    ),
                },
                assumptions=_unique_strings([str(item) for item in source.get("assumptions", []) if str(item).strip()]),
                workbook_mapping=source.get("workbook_mapping", {}) if isinstance(source.get("workbook_mapping", {}), dict) else {},
                service_tags=_unique_strings([str(item) for item in source.get("service_tags", []) if str(item).strip()]),
                region_tags=_unique_strings([str(item) for item in source.get("region_tags", []) if str(item).strip()]),
            )
        return canonical, ValidationReportArtifact(
            status="VALID" if not errors else "INVALID",
            warnings=_unique_strings(warnings),
            errors=_unique_strings(errors),
            details={"required_anchor_count": len(REQUIRED_PROPOSAL_TEMPLATE_ANCHORS)},
        )

    def _canonical_from_rows(
        self,
        rows: list[dict[str, Any]],
        *,
        sheet_name: str | None = None,
        header_row: int | None = None,
    ) -> tuple[ClientPricingModelCanonicalArtifact | None, ValidationReportArtifact]:
        errors: list[str] = []
        warnings: list[str] = []
        normalized_rows: list[PricingModelLaborCategoryArtifact] = []
        for item in rows:
            code = _normalize_key(str(item.get("code") or item.get("labor_category") or item.get("label") or ""))
            if not code:
                continue
            label = _normalize_space(str(item.get("label") or item.get("name") or code.replace("_", " ").title()))
            try:
                hourly_rate = float(item.get("hourly_rate", 0))
                burden_factor = float(item.get("burden_factor", 0))
                markup_factor = float(item.get("markup_factor", 0))
                default_hours = float(item.get("default_hours_per_week", _default_hours_for_code(code)))
            except (TypeError, ValueError):
                errors.append(f"Pricing row '{label or code}' contains invalid numeric values.")
                continue
            if "default_hours_per_week" not in item:
                warnings.append(f"'{label}' did not include default hours; a standard workload assumption was applied.")
            normalized_rows.append(
                PricingModelLaborCategoryArtifact(
                    code=code,
                    label=label or code.replace("_", " ").title(),
                    hourly_rate=hourly_rate,
                    burden_factor=burden_factor,
                    markup_factor=markup_factor,
                    default_hours_per_week=default_hours,
                    apply_site_multiplier=code not in {"project_manager"},
                    apply_continuous_coverage_multiplier=code in {"janitor", "day_porter"},
                )
            )
        if not normalized_rows:
            errors.append("No pricing rows matched the expected workbook or CSV headers.")
        canonical = None
        if not errors:
            canonical = ClientPricingModelCanonicalArtifact(
                labor_categories=normalized_rows,
                staffing_rules={"site_multiplier": 1.0, "continuous_coverage_multiplier": 1.35},
                assumptions=[
                    "Imported pricing models preserve the client source file and normalize it into the internal pricing schema.",
                    "Default hours per week should be confirmed during implementation and pricing QA.",
                ],
                workbook_mapping={
                    "sheet_name": sheet_name or "Pricing Summary",
                    "header_row": header_row or 1,
                    "column_headers": {
                        "labor_category": "Labor Category",
                        "hours_per_week": "Hours/Week",
                        "hourly_rate": "Hourly Rate",
                        "burden_factor": "Burden",
                        "markup_factor": "Markup",
                        "loaded_hourly_cost": "Loaded Hourly Cost",
                        "annual_sell_price": "Annual Sell Price",
                    },
                },
            )
        return canonical, ValidationReportArtifact(
            status="VALID" if not errors else "INVALID",
            warnings=_unique_strings(warnings),
            errors=_unique_strings(errors),
            details={"sheet_name": sheet_name, "header_row": header_row},
        )

    def _read_csv_model(self, content: bytes) -> tuple[ClientPricingModelCanonicalArtifact | None, ValidationReportArtifact]:
        decoded = io.StringIO(content.decode("utf-8-sig"))
        reader = csv.DictReader(decoded)
        rows: list[dict[str, Any]] = []
        for raw_row in reader:
            rows.append({_header_key(key): value for key, value in raw_row.items() if key})
        return self._canonical_from_rows(rows, sheet_name="Imported CSV", header_row=1)

    def _read_xlsx_model(self, content: bytes) -> tuple[ClientPricingModelCanonicalArtifact | None, ValidationReportArtifact]:
        workbook = load_workbook(io.BytesIO(content))
        for worksheet in workbook.worksheets:
            for row_index, row in enumerate(worksheet.iter_rows(min_row=1, max_row=5, values_only=True), start=1):
                headers = [_header_key(str(cell or "")) for cell in row]
                if not {"hourly_rate", "burden_factor", "markup_factor"} <= set(headers):
                    continue
                parsed_rows: list[dict[str, Any]] = []
                for data_row in worksheet.iter_rows(min_row=row_index + 1, values_only=True):
                    row_payload = {
                        headers[index]: data_row[index]
                        for index in range(min(len(headers), len(data_row)))
                        if headers[index]
                    }
                    if not any(str(value).strip() for value in row_payload.values()):
                        continue
                    parsed_rows.append(row_payload)
                return self._canonical_from_rows(parsed_rows, sheet_name=worksheet.title, header_row=row_index)
        return None, ValidationReportArtifact(
            status="INVALID",
            warnings=[],
            errors=[
                "The uploaded workbook does not match the supported import layout. "
                "Use a sheet with headers for labor category, hourly rate, burden factor, and markup factor."
            ],
            details={},
        )

    def _store_uploaded_file(self, directory: Path, filename: str, content: bytes) -> Path:
        extension = Path(filename).suffix.lower()
        target = directory / f"{uuid.uuid4().hex}-{_safe_slug(Path(filename).stem)}{extension}"
        target.write_bytes(content)
        return target

    def import_pricing_model(
        self,
        profile_id: str,
        *,
        name: str | None,
        filename: str,
        content: bytes,
    ) -> ClientPricingModelResponse:
        profile = self._profile(profile_id)
        extension = Path(filename).suffix.lower()
        if extension not in {".json", ".csv", ".xlsx"}:
            raise ClientProfileError("Pricing model import supports .json, .csv, and .xlsx files.")

        if extension == ".json":
            canonical, report = self._canonical_from_json(json.loads(content.decode("utf-8")))
        elif extension == ".csv":
            canonical, report = self._read_csv_model(content)
        else:
            canonical, report = self._read_xlsx_model(content)

        model_storage = _storage_dir_for_profile(profile_id, "pricing-models")
        source_path = self._store_uploaded_file(model_storage, filename, content)
        pricing_model = ClientPricingModel(
            client_profile_id=profile.id,
            name=_normalize_space(name or Path(filename).stem.replace("-", " ").replace("_", " ")),
            status="APPROVED" if report.status == "VALID" else "DRAFT",
            onboarding_mode="IMPORT",
            source_format=extension.lstrip("."),
            source_filename=filename,
            source_file_path=source_path.as_posix(),
            workbook_template_filename=filename if extension == ".xlsx" else None,
            workbook_template_path=source_path.as_posix() if extension == ".xlsx" and report.status == "VALID" else None,
            canonical_model_json=_dump_json(canonical.model_dump()) if canonical else None,
            validation_report_json=_dump_json(report.model_dump()),
        )
        self.db.add(pricing_model)
        self.db.flush()
        if not profile.default_pricing_model_id and pricing_model.status == "APPROVED":
            profile.default_pricing_model_id = pricing_model.id
        log_audit_event(
            self.db,
            opportunity_id=None,
            actor="admin",
            action="client_pricing_model_imported",
            after_state_json=_dump_json({"client_profile_id": profile.id, "pricing_model_id": pricing_model.id}),
        )
        self.db.commit()
        self.db.refresh(pricing_model)
        return _serialize_pricing_model(pricing_model)

    def create_pricing_model(self, profile_id: str, payload: ClientPricingModelCreateRequest) -> ClientPricingModelResponse:
        profile = self._profile(profile_id)
        if not payload.labor_categories:
            raise ClientProfileError("At least one labor category is required to build a pricing model.")
        canonical = ClientPricingModelCanonicalArtifact(
            labor_categories=payload.labor_categories,
            staffing_rules={
                "site_multiplier": payload.site_multiplier,
                "continuous_coverage_multiplier": payload.continuous_coverage_multiplier,
            },
            assumptions=_unique_strings(payload.assumptions),
            workbook_mapping={},
            service_tags=_unique_strings(payload.service_tags),
            region_tags=_unique_strings(payload.region_tags),
        )
        report = ValidationReportArtifact(status="VALID", warnings=[], errors=[], details={})
        pricing_model = ClientPricingModel(
            client_profile_id=profile.id,
            name=_normalize_space(payload.name),
            status="APPROVED",
            onboarding_mode="BUILD",
            source_format="builder",
            canonical_model_json=_dump_json(canonical.model_dump()),
            validation_report_json=_dump_json(report.model_dump()),
        )
        self.db.add(pricing_model)
        self.db.flush()
        if not profile.default_pricing_model_id:
            profile.default_pricing_model_id = pricing_model.id
        log_audit_event(
            self.db,
            opportunity_id=None,
            actor="admin",
            action="client_pricing_model_created",
            after_state_json=_dump_json({"client_profile_id": profile.id, "pricing_model_id": pricing_model.id}),
        )
        self.db.commit()
        self.db.refresh(pricing_model)
        return _serialize_pricing_model(pricing_model)

    def upload_proposal_template(self, profile_id: str, *, name: str | None, filename: str, content: bytes) -> ClientProposalTemplateResponse:
        profile = self._profile(profile_id)
        if Path(filename).suffix.lower() != ".docx":
            raise ClientProfileError("Proposal templates must be uploaded as .docx files.")
        template_storage = _storage_dir_for_profile(profile_id, "proposal-templates")
        template_path = self._store_uploaded_file(template_storage, filename, content)
        report = ValidationReportArtifact.model_validate(validate_proposal_template(template_path))
        template = ClientProposalTemplate(
            client_profile_id=profile.id,
            name=_normalize_space(name or Path(filename).stem.replace("-", " ").replace("_", " ")),
            status="APPROVED" if not report.errors else "DRAFT",
            source_filename=filename,
            storage_path=template_path.as_posix(),
            anchor_map_json=_dump_json(report.details.get("anchors", {})),
            validation_report_json=_dump_json(report.model_dump()),
        )
        self.db.add(template)
        self.db.flush()
        if not profile.default_proposal_template_id and template.status == "APPROVED":
            profile.default_proposal_template_id = template.id
        log_audit_event(
            self.db,
            opportunity_id=None,
            actor="admin",
            action="client_proposal_template_uploaded",
            after_state_json=_dump_json({"client_profile_id": profile.id, "proposal_template_id": template.id}),
        )
        self.db.commit()
        self.db.refresh(template)
        return _serialize_template(template)

    def set_defaults(self, profile_id: str, payload: ClientProfileDefaultsRequest) -> ClientProfileResponse:
        profile = self._profile(profile_id)
        if payload.default_pricing_model_id:
            model = self._pricing_model(payload.default_pricing_model_id)
            if model.client_profile_id != profile.id or model.status != "APPROVED":
                raise ClientProfileError("Default pricing model must belong to the profile and be approved.")
            profile.default_pricing_model_id = model.id
        else:
            profile.default_pricing_model_id = None
        if payload.default_proposal_template_id:
            template = self._proposal_template(payload.default_proposal_template_id)
            if template.client_profile_id != profile.id or template.status != "APPROVED":
                raise ClientProfileError("Default proposal template must belong to the profile and be approved.")
            profile.default_proposal_template_id = template.id
        else:
            profile.default_proposal_template_id = None
        self.db.commit()
        self.db.refresh(profile)
        return self._serialize_profile(profile, include_assets=True)

    def _ensure_config_row(self, opportunity: Opportunity) -> OpportunityProposalConfig:
        existing = self._config_row(opportunity.id)
        if existing:
            return existing
        created = OpportunityProposalConfig(opportunity_id=opportunity.id)
        self.db.add(created)
        self.db.flush()
        return created

    def _auto_match_profile(self, opportunity: Opportunity) -> ClientProfile | None:
        active_profiles = [profile for profile in self._list_profiles() if profile.status.upper() == "ACTIVE"]
        if len(active_profiles) == 1:
            return active_profiles[0]
        for profile in active_profiles:
            if profile.status.upper() != "ACTIVE":
                continue
            if self._profile_matches_name(profile, opportunity.client):
                return profile
        return None

    def _config_validation(
        self,
        *,
        profile: ClientProfile | None,
        pricing_model: ClientPricingModel | None,
        proposal_template: ClientProposalTemplate | None,
    ) -> tuple[list[str], list[str]]:
        warnings: list[str] = []
        errors: list[str] = []
        if profile is None:
            warnings.append("No client profile is selected. The app will use the standard pricing/export fallback for internal workflows.")
            return _unique_strings(warnings), _unique_strings(errors)
        if profile.client_environment_status == "SETUP_GAPS":
            warnings.append("Client onboarding is active, but setup gaps still remain for this environment.")
        if pricing_model is None:
            errors.append("No approved client pricing model is selected for this opportunity.")
        else:
            report = _validation_report(_parse_json(pricing_model.validation_report_json, {}))
            if pricing_model.status != "APPROVED" or (report and report.errors):
                errors.append("The selected pricing model is not approved for runtime use.")
        if proposal_template is None:
            errors.append("No approved client proposal template is selected for this opportunity.")
        else:
            report = _validation_report(_parse_json(proposal_template.validation_report_json, {}))
            if proposal_template.status != "APPROVED" or (report and report.errors):
                errors.append("The selected proposal template failed validation and cannot be used for export.")
        return _unique_strings(warnings), _unique_strings(errors)

    def resolve_opportunity_config(self, opportunity_id: str) -> ResolvedOpportunityConfig:
        opportunity = self._opportunity(opportunity_id)
        config_row = self._ensure_config_row(opportunity)
        profile = self.db.get(ClientProfile, config_row.client_profile_id) if config_row.client_profile_id else None
        if profile is None:
            matched = self._auto_match_profile(opportunity)
            if matched is not None:
                profile = matched
                config_row.client_profile_id = matched.id
                config_row.client_profile_selection_source = "default"

        if profile is not None:
            if (
                config_row.pricing_model_selection_source != "override"
                and profile.default_pricing_model_id
                and config_row.pricing_model_id != profile.default_pricing_model_id
            ):
                config_row.pricing_model_id = profile.default_pricing_model_id
                config_row.pricing_model_selection_source = "default"
            if (
                config_row.proposal_template_selection_source != "override"
                and profile.default_proposal_template_id
                and config_row.proposal_template_id != profile.default_proposal_template_id
            ):
                config_row.proposal_template_id = profile.default_proposal_template_id
                config_row.proposal_template_selection_source = "default"

        pricing_model = self.db.get(ClientPricingModel, config_row.pricing_model_id) if config_row.pricing_model_id else None
        proposal_template = self.db.get(ClientProposalTemplate, config_row.proposal_template_id) if config_row.proposal_template_id else None
        warnings, errors = self._config_validation(
            profile=profile,
            pricing_model=pricing_model,
            proposal_template=proposal_template,
        )
        config_row.validation_warnings_json = _dump_json(warnings)
        config_row.validation_errors_json = _dump_json(errors)
        self.db.flush()
        config = OpportunityProposalConfigArtifact(
            client_profile_id=profile.id if profile else None,
            client_profile_name=profile.display_name if profile else None,
            active_onboarding_pack_id=profile.active_onboarding_pack_id if profile else None,
            active_onboarding_pack_name=(
                self.db.get(ClientOnboardingPack, profile.active_onboarding_pack_id).pack_name
                if profile and profile.active_onboarding_pack_id and self.db.get(ClientOnboardingPack, profile.active_onboarding_pack_id)
                else None
            ),
            active_playbook_version=(
                self.db.get(ClientPlaybook, profile.active_playbook_id).version
                if profile and profile.active_playbook_id and self.db.get(ClientPlaybook, profile.active_playbook_id)
                else None
            ),
            client_environment_status=profile.client_environment_status if profile else None,
            pricing_model_id=pricing_model.id if pricing_model else None,
            pricing_model_name=pricing_model.name if pricing_model else None,
            proposal_template_id=proposal_template.id if proposal_template else None,
            proposal_template_name=proposal_template.name if proposal_template else None,
            client_profile_selection_source=config_row.client_profile_selection_source,
            pricing_model_selection_source=config_row.pricing_model_selection_source,
            proposal_template_selection_source=config_row.proposal_template_selection_source,
            setup_gaps=(
                _parse_json(self.db.get(ClientOnboardingPack, profile.active_onboarding_pack_id).setup_gaps_json, [])
                if profile and profile.active_onboarding_pack_id and self.db.get(ClientOnboardingPack, profile.active_onboarding_pack_id)
                else []
            ),
            validation_warnings=warnings,
            validation_errors=errors,
        )
        return ResolvedOpportunityConfig(
            config_row=config_row,
            config=config,
            available_profiles=self._profile_options(),
            available_pricing_models=self._pricing_model_options(profile.id if profile else None),
            available_proposal_templates=self._proposal_template_options(profile.id if profile else None),
        )

    def opportunity_config(self, opportunity_id: str) -> OpportunityProposalConfigArtifact:
        return self.resolve_opportunity_config(opportunity_id).config

    def selected_pricing_model(self, opportunity_id: str) -> ClientPricingModelResponse | None:
        resolved = self.resolve_opportunity_config(opportunity_id)
        if not resolved.config.pricing_model_id:
            return None
        return _serialize_pricing_model(self._pricing_model(resolved.config.pricing_model_id))

    def selected_proposal_template(self, opportunity_id: str) -> ClientProposalTemplateResponse | None:
        resolved = self.resolve_opportunity_config(opportunity_id)
        if not resolved.config.proposal_template_id:
            return None
        return _serialize_template(self._proposal_template(resolved.config.proposal_template_id))

    def apply_snapshot_to_run(self, opportunity_id: str, run: Any) -> OpportunityProposalConfigArtifact:
        resolved = self.resolve_opportunity_config(opportunity_id)
        run.client_profile_id = resolved.config.client_profile_id
        run.onboarding_pack_id = resolved.config.active_onboarding_pack_id
        run.playbook_id = self.db.get(ClientProfile, resolved.config.client_profile_id).active_playbook_id if resolved.config.client_profile_id else None
        run.playbook_version = resolved.config.active_playbook_version
        run.pricing_model_id = resolved.config.pricing_model_id
        run.proposal_template_id = resolved.config.proposal_template_id
        run.selection_source_json = _dump_json(
            {
                "client_profile": resolved.config.client_profile_selection_source,
                "onboarding_pack": "active_environment" if resolved.config.active_onboarding_pack_id else "missing",
                "pricing_model": resolved.config.pricing_model_selection_source,
                "proposal_template": resolved.config.proposal_template_selection_source,
            }
        )
        return resolved.config

    def select_profile(self, opportunity_id: str, payload: OpportunityProposalConfigSelectionRequest) -> OpportunityProposalConfigArtifact:
        opportunity = self._opportunity(opportunity_id)
        config_row = self._ensure_config_row(opportunity)
        if payload.target_id:
            profile = self._profile(payload.target_id)
            config_row.client_profile_id = profile.id
            config_row.client_profile_selection_source = "override"
            config_row.pricing_model_id = profile.default_pricing_model_id
            config_row.pricing_model_selection_source = "default" if profile.default_pricing_model_id else "missing"
            config_row.proposal_template_id = profile.default_proposal_template_id
            config_row.proposal_template_selection_source = "default" if profile.default_proposal_template_id else "missing"
        else:
            config_row.client_profile_id = None
            config_row.client_profile_selection_source = "missing"
            config_row.pricing_model_id = None
            config_row.pricing_model_selection_source = "missing"
            config_row.proposal_template_id = None
            config_row.proposal_template_selection_source = "missing"
        log_audit_event(
            self.db,
            opportunity_id=opportunity_id,
            actor=payload.actor,
            action="opportunity_proposal_config_profile_selected",
            after_state_json=_dump_json({"target_id": payload.target_id}),
        )
        resolved = self.resolve_opportunity_config(opportunity_id)
        self.db.commit()
        return resolved.config

    def select_pricing_model(self, opportunity_id: str, payload: OpportunityProposalConfigSelectionRequest) -> OpportunityProposalConfigArtifact:
        opportunity = self._opportunity(opportunity_id)
        config_row = self._ensure_config_row(opportunity)
        profile = self.db.get(ClientProfile, config_row.client_profile_id) if config_row.client_profile_id else None
        if payload.target_id:
            model = self._pricing_model(payload.target_id)
            if profile and model.client_profile_id != profile.id:
                raise ClientProfileError("Pricing model must belong to the selected client profile.")
            config_row.pricing_model_id = model.id
            config_row.pricing_model_selection_source = "override"
        else:
            config_row.pricing_model_id = profile.default_pricing_model_id if profile and profile.default_pricing_model_id else None
            config_row.pricing_model_selection_source = "default" if config_row.pricing_model_id else "missing"
        log_audit_event(
            self.db,
            opportunity_id=opportunity_id,
            actor=payload.actor,
            action="opportunity_proposal_config_pricing_selected",
            after_state_json=_dump_json({"target_id": payload.target_id}),
        )
        resolved = self.resolve_opportunity_config(opportunity_id)
        self.db.commit()
        return resolved.config

    def select_template(self, opportunity_id: str, payload: OpportunityProposalConfigSelectionRequest) -> OpportunityProposalConfigArtifact:
        opportunity = self._opportunity(opportunity_id)
        config_row = self._ensure_config_row(opportunity)
        profile = self.db.get(ClientProfile, config_row.client_profile_id) if config_row.client_profile_id else None
        if payload.target_id:
            template = self._proposal_template(payload.target_id)
            if profile and template.client_profile_id != profile.id:
                raise ClientProfileError("Proposal template must belong to the selected client profile.")
            config_row.proposal_template_id = template.id
            config_row.proposal_template_selection_source = "override"
        else:
            config_row.proposal_template_id = profile.default_proposal_template_id if profile and profile.default_proposal_template_id else None
            config_row.proposal_template_selection_source = "default" if config_row.proposal_template_id else "missing"
        log_audit_event(
            self.db,
            opportunity_id=opportunity_id,
            actor=payload.actor,
            action="opportunity_proposal_config_template_selected",
            after_state_json=_dump_json({"target_id": payload.target_id}),
        )
        resolved = self.resolve_opportunity_config(opportunity_id)
        self.db.commit()
        return resolved.config
