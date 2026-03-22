from __future__ import annotations

import csv
import zipfile
from pathlib import Path

from docx import Document

from app.modules.proposal_builder.content_authority import load_form_templates, select_assets
from app.modules.proposal_builder.schemas import FormPackageArtifact, GeneratedFileArtifact, PricingPackageArtifact


def _write_docx(path: Path, heading: str, paragraphs: list[str]) -> None:
    document = Document()
    document.add_heading(heading, level=0)
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    document.save(path)


def _write_markdown(path: Path, heading: str, bullets: list[str]) -> None:
    lines = [f"# {heading}", ""]
    lines.extend(f"- {item}" for item in bullets)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def build_form_package(
    *,
    workspace: dict[str, object],
    pricing_package: PricingPackageArtifact,
    target_dir: Path,
    approved_assets: list[dict[str, object]] | None = None,
    form_templates: list[dict[str, object]] | None = None,
) -> FormPackageArtifact:
    target_dir.mkdir(parents=True, exist_ok=True)
    templates = form_templates if form_templates is not None else load_form_templates()
    text_blob = " ".join(
        [
            str(workspace.get("client_name", "")),
            str(workspace.get("opportunity_name", "")),
            " ".join(
                str(item.get("requirement_text", ""))
                for item in workspace.get("requirements_list", [])
                if isinstance(item, dict)
            ),
        ]
    ).lower()
    tags = ["janitorial"]
    if "airport" in text_blob:
        tags.extend(["airport", "public-sector"])

    completed_forms: list[GeneratedFileArtifact] = []
    checklist: list[str] = []
    unresolved_field_map_gaps: list[str] = []
    blocked_forms: list[str] = []

    transmittal_path = target_dir / "transmittal-letter.docx"
    _write_docx(
        transmittal_path,
        "Proposal Transmittal Letter",
        [
            f"Buyer: {workspace.get('client_name', '')}",
            f"Opportunity: {workspace.get('opportunity_name', '')}",
            "This transmittal package includes the narrative draft, pricing workbook, attachment checklist, and supporting packets.",
        ],
    )
    completed_forms.append(
        GeneratedFileArtifact(
            label="Transmittal Letter",
            path=transmittal_path.as_posix(),
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    )
    checklist.append("Review and sign the transmittal letter before final production.")

    if "addenda" in text_blob:
        addenda_path = target_dir / "addenda-acknowledgement.docx"
        _write_docx(
            addenda_path,
            "Addenda Acknowledgement",
            [
                f"Opportunity: {workspace.get('opportunity_name', '')}",
                "Confirm each addendum number and receipt date before submission.",
            ],
        )
        completed_forms.append(
            GeneratedFileArtifact(
                label="Addenda Acknowledgement",
                path=addenda_path.as_posix(),
                media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        )
        checklist.append("Validate addenda numbers against the final solicitation package.")

    resume_assets = select_assets(asset_type="resume", tags=tags, assets=approved_assets)
    reference_assets = select_assets(asset_type="reference", tags=tags, assets=approved_assets)
    insurance_assets = select_assets(asset_type="insurance", tags=tags, assets=approved_assets)
    if not resume_assets:
        blocked_forms.append("Missing approved resume packet content.")
    if not reference_assets:
        blocked_forms.append("Missing approved reference packet content.")
    if not insurance_assets:
        blocked_forms.append("Missing current insurance summary.")

    resumes_packet = target_dir / "resume-packet.md"
    _write_markdown(resumes_packet, "Resume Packet", [asset.body for asset in resume_assets] or ["No approved resume assets available."])
    references_packet = target_dir / "reference-packet.md"
    _write_markdown(references_packet, "Reference Packet", [asset.body for asset in reference_assets] or ["No approved references available."])
    insurance_packet = target_dir / "insurance-summary.md"
    _write_markdown(insurance_packet, "Insurance Summary", [asset.body for asset in insurance_assets] or ["No current insurance summary available."])

    checklist_csv = target_dir / "attachment-checklist.csv"
    with checklist_csv.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["Attachment", "Status"])
        writer.writerow(["Pricing Workbook", "Ready" if pricing_package.workbook_path else "Blocked"])
        writer.writerow(["Resume Packet", "Ready" if resume_assets else "Blocked"])
        writer.writerow(["Reference Packet", "Ready" if reference_assets else "Blocked"])
        writer.writerow(["Insurance Summary", "Ready" if insurance_assets else "Blocked"])

    bundle_path = target_dir / "completed-forms-bundle.zip"
    with zipfile.ZipFile(bundle_path, "w", compression=zipfile.ZIP_DEFLATED) as bundle:
        for path in (transmittal_path, resumes_packet, references_packet, insurance_packet, checklist_csv):
            if path.exists():
                bundle.write(path, arcname=path.name)

    completed_forms.append(GeneratedFileArtifact(label="Completed Forms Bundle", path=bundle_path.as_posix(), media_type="application/zip"))
    checklist.extend(
        [
            "Confirm all required forms listed in the RFP are either completed or flagged for manual production.",
            "Confirm pricing workbook version and filename match the solicitation instructions.",
        ]
    )
    if not any(template.get("id") == "default-pricing-workbook" for template in templates):
        unresolved_field_map_gaps.append("No default pricing workbook template metadata is available.")

    return FormPackageArtifact(
        completed_forms=completed_forms,
        attachment_checklist=checklist,
        resumes_packet_paths=[resumes_packet.as_posix()],
        references_packet_paths=[references_packet.as_posix()],
        unresolved_field_map_gaps=unresolved_field_map_gaps,
        blocked_forms=blocked_forms,
    )
