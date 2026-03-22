from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import _Cell
from docx.text.paragraph import Paragraph


ANCHOR_PATTERN = re.compile(r"\{\{\s*([a-z_]+)\s*\}\}")
REQUIRED_PROPOSAL_TEMPLATE_ANCHORS = (
    "transmittal_letter",
    "executive_summary",
    "technical_approach",
    "management_staffing_plan",
    "transition_mobilization_plan",
    "past_performance",
    "pricing_and_commercials",
    "compliance_attachments",
)
SECTION_TITLE_BY_ANCHOR = {
    "executive_summary": "Executive Summary",
    "technical_approach": "Technical Approach",
    "management_staffing_plan": "Management and Staffing Plan",
    "transition_mobilization_plan": "Transition and Mobilization Plan",
    "past_performance": "Past Performance",
    "pricing_and_commercials": "Pricing and Commercials",
    "compliance_attachments": "Compliance and Attachments",
}
METADATA_PLACEHOLDERS = (
    "opportunity_name",
    "client_name",
    "solicitation_number",
    "proposal_due",
)


def _iter_paragraphs(document: Document) -> list[Paragraph]:
    paragraphs = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    return paragraphs


def _iter_cells(document: Document) -> list[_Cell]:
    cells: list[_Cell] = []
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                cells.append(cell)
    return cells


def _iter_story_parts(document: Document) -> list[Any]:
    story_parts: list[Any] = [document]
    for section in document.sections:
        story_parts.append(section.header)
        story_parts.append(section.footer)
    return story_parts


def _iter_story_paragraphs(document: Document) -> list[Paragraph]:
    paragraphs: list[Paragraph] = []
    for story in _iter_story_parts(document):
        paragraphs.extend(story.paragraphs)
        for table in story.tables:
            for row in table.rows:
                for cell in row.cells:
                    paragraphs.extend(cell.paragraphs)
    return paragraphs


def _iter_story_cells(document: Document) -> list[_Cell]:
    cells: list[_Cell] = []
    for story in _iter_story_parts(document):
        for table in story.tables:
            for row in table.rows:
                for cell in row.cells:
                    cells.append(cell)
    return cells


def _replace_plain_placeholder(text: str, replacements: dict[str, str]) -> str:
    updated = text
    for key, value in replacements.items():
        updated = updated.replace(f"{{{{ {key} }}}}", value)
        updated = updated.replace(f"{{{{{key}}}}}", value)
    return updated


def _replace_plain_placeholders_in_paragraph(paragraph: Paragraph, replacements: dict[str, str]) -> None:
    if paragraph.runs:
        for run in paragraph.runs:
            run.text = _replace_plain_placeholder(run.text or "", replacements)
        return
    paragraph.text = _replace_plain_placeholder(paragraph.text or "", replacements)


def validate_proposal_template(path: Path) -> dict[str, Any]:
    document = Document(path)
    anchor_locations: dict[str, list[str]] = {}
    for index, paragraph in enumerate(_iter_paragraphs(document), start=1):
        for match in ANCHOR_PATTERN.findall(paragraph.text or ""):
            anchor_locations.setdefault(match, []).append(f"paragraph:{index}")
    errors: list[str] = []
    warnings: list[str] = []
    missing = [anchor for anchor in REQUIRED_PROPOSAL_TEMPLATE_ANCHORS if anchor not in anchor_locations]
    duplicates = [anchor for anchor, locations in anchor_locations.items() if anchor in REQUIRED_PROPOSAL_TEMPLATE_ANCHORS and len(locations) > 1]
    if missing:
        errors.append("Missing required template anchors: " + ", ".join(missing))
    if duplicates:
        errors.append("Duplicate required template anchors: " + ", ".join(duplicates))
    unsupported = sorted(set(anchor_locations) - set(REQUIRED_PROPOSAL_TEMPLATE_ANCHORS) - set(METADATA_PLACEHOLDERS))
    if unsupported:
        warnings.append("Template contains unrecognized anchors: " + ", ".join(unsupported))
    return {
        "status": "VALID" if not errors else "INVALID",
        "warnings": warnings,
        "errors": errors,
        "details": {
            "anchors": anchor_locations,
            "missing_anchors": missing,
            "duplicate_anchors": duplicates,
            "unsupported_anchors": unsupported,
        },
    }


def _insert_paragraph_after(paragraph: Paragraph, text: str = "", style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_paragraph = Paragraph(new_p, paragraph._parent)
    if style:
        try:
            new_paragraph.style = style
        except Exception:
            pass
    if text:
        new_paragraph.add_run(text)
    return new_paragraph


def _append_page_number_run(paragraph: Paragraph, *, prefix: str = "") -> None:
    if prefix:
        paragraph.add_run(prefix)
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(text)
    run._r.append(end)


def _render_markdownish_into_document(anchor_paragraph: Paragraph, body: str, *, section_title: str | None = None) -> None:
    current = anchor_paragraph
    if section_title:
        current.text = section_title
        try:
            current.style = "Heading 1"
        except Exception:
            pass
        current = _insert_paragraph_after(current)
        current.text = ""
    else:
        current.text = ""
        current.style = "Normal"
    first_written = False
    paragraph_lines: list[str] = []

    def flush_paragraph() -> None:
        nonlocal current, first_written
        if not paragraph_lines:
            return
        target = current if not first_written else _insert_paragraph_after(current)
        target.text = " ".join(paragraph_lines).strip()
        current = target
        first_written = True
        paragraph_lines.clear()

    for raw_line in body.splitlines():
        stripped = raw_line.strip()
        if not stripped:
            flush_paragraph()
            continue
        if stripped.startswith("### "):
            flush_paragraph()
            target = current if not first_written else _insert_paragraph_after(current)
            target.text = stripped[4:].strip()
            try:
                target.style = "Heading 2"
            except Exception:
                pass
            current = target
            first_written = True
            continue
        if stripped.startswith("- "):
            flush_paragraph()
            target = current if not first_written else _insert_paragraph_after(current)
            target.text = stripped[2:].strip()
            try:
                target.style = "List Bullet"
            except Exception:
                pass
            current = target
            first_written = True
            continue
        paragraph_lines.append(stripped)

    flush_paragraph()


def _section_body_map(workspace: dict[str, Any]) -> dict[str, str]:
    sections = workspace.get("client_sections") or (workspace.get("draft_package") or {}).get("client_sections") or []
    title_map = {str(section.get("section_title", "")).lower(): str(section.get("body_markdown", "")).strip() for section in sections}
    return {
        "transmittal_letter": title_map.get("transmittal letter", ""),
        "executive_summary": title_map.get("executive summary", ""),
        "technical_approach": title_map.get("technical approach", ""),
        "management_staffing_plan": title_map.get("management and staffing plan", ""),
        "transition_mobilization_plan": title_map.get("transition and mobilization plan", ""),
        "past_performance": title_map.get("past performance", ""),
        "pricing_and_commercials": title_map.get("pricing and commercials", ""),
        "compliance_attachments": title_map.get("compliance and attachments", ""),
    }


def render_proposal_template(*, workspace: dict[str, Any], template_path: Path, output_path: Path) -> None:
    document = Document(template_path)
    summary = workspace.get("opportunity_summary") or {}
    replacements = {
        "opportunity_name": str(summary.get("opportunity_name") or workspace.get("opportunity_name") or ""),
        "client_name": str(summary.get("client_name") or workspace.get("client_name") or ""),
        "solicitation_number": str(summary.get("solicitation_number") or workspace.get("solicitation_id") or ""),
        "proposal_due": " ".join(
            part
            for part in [
                str(summary.get("proposal_due_date") or workspace.get("extracted_deadline") or "").strip(),
                str(summary.get("proposal_due_time") or "").strip(),
            ]
            if part
        ),
    }
    section_bodies = _section_body_map(workspace)

    for paragraph in _iter_story_paragraphs(document):
        _replace_plain_placeholders_in_paragraph(paragraph, replacements)

    for paragraph in _iter_paragraphs(document):
        matches = ANCHOR_PATTERN.findall(paragraph.text or "")
        if not matches:
            continue
        anchor = matches[0]
        if anchor in section_bodies:
            _render_markdownish_into_document(paragraph, section_bodies[anchor], section_title=SECTION_TITLE_BY_ANCHOR.get(anchor))
        elif anchor in replacements:
            paragraph.text = _replace_plain_placeholder(paragraph.text or "", replacements)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
