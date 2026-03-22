from __future__ import annotations

from copy import deepcopy
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

from app.modules.proposal_builder.content_authority import select_assets
from app.modules.proposal_builder.schemas import (
    ClientPricingModelCanonicalArtifact,
    ClientPricingModelResponse,
    PricingLineItemArtifact,
    PricingPackageArtifact,
)


def _workspace_text(workspace: dict[str, object]) -> str:
    parts: list[str] = []
    summary = workspace.get("opportunity_summary") or {}
    structured_fields = workspace.get("structured_fields") or {}
    parts.extend(str(value) for value in summary.values() if isinstance(value, str))
    parts.extend(str(value) for value in structured_fields.values() if isinstance(value, str))
    for requirement in workspace.get("requirements_list", []):
        if isinstance(requirement, dict):
            parts.append(str(requirement.get("requirement_text", "")))
    return " ".join(parts).lower()


def _site_count(workspace: dict[str, object]) -> int:
    structured_fields = workspace.get("structured_fields") or {}
    sites = structured_fields.get("site_geography") or []
    return max(len(sites), 1)


def _service_tags(workspace: dict[str, object]) -> list[str]:
    text = _workspace_text(workspace)
    tags = {"janitorial"}
    if "airport" in text:
        tags.add("airport")
    if "day porter" in text or "day-porter" in text:
        tags.add("day-porter")
    if "public" in text or "authority" in text or "agency" in text:
        tags.add("public-sector")
    return sorted(tags)


def _staffing_template(workspace: dict[str, object]) -> list[tuple[str, float]]:
    text = _workspace_text(workspace)
    sites = _site_count(workspace)
    janitor_hours = 160.0 * sites
    porter_hours = 80.0 * sites if "day porter" in text or "day-porter" in text else 40.0 * sites
    supervisor_hours = 20.0 * sites
    project_manager_hours = 6.0 * sites
    if "24/7" in text or "seven days" in text or "7 days" in text:
        janitor_hours *= 1.35
        porter_hours *= 1.25
    return [
        ("janitor", round(janitor_hours, 2)),
        ("day_porter", round(porter_hours, 2)),
        ("supervisor", round(supervisor_hours, 2)),
        ("project_manager", round(project_manager_hours, 2)),
    ]


def _write_template_workbook(
    *,
    template_path: Path,
    target_path: Path,
    line_items: list[PricingLineItemArtifact],
    mapping: dict[str, Any],
) -> None:
    workbook = load_workbook(template_path)
    sheet_name = str(mapping.get("sheet_name") or workbook.sheetnames[0])
    worksheet = workbook[sheet_name] if sheet_name in workbook.sheetnames else workbook[workbook.sheetnames[0]]
    header_row = int(mapping.get("header_row") or 1)
    headers = {str(cell.value or "").strip(): index + 1 for index, cell in enumerate(worksheet[header_row])}
    column_headers = mapping.get("column_headers", {})
    row_cursor = header_row + 1
    for item in line_items:
        worksheet.cell(row=row_cursor, column=headers.get(column_headers.get("labor_category", "Labor Category"), 1), value=item.labor_category)
        worksheet.cell(row=row_cursor, column=headers.get(column_headers.get("hours_per_week", "Hours/Week"), 2), value=item.hours_per_week)
        worksheet.cell(row=row_cursor, column=headers.get(column_headers.get("hourly_rate", "Hourly Rate"), 3), value=item.hourly_rate)
        worksheet.cell(row=row_cursor, column=headers.get(column_headers.get("burden_factor", "Burden"), 4), value=item.burden_factor)
        worksheet.cell(row=row_cursor, column=headers.get(column_headers.get("markup_factor", "Markup"), 5), value=item.markup_factor)
        worksheet.cell(
            row=row_cursor,
            column=headers.get(column_headers.get("loaded_hourly_cost", "Loaded Hourly Cost"), 6),
            value=item.loaded_hourly_cost,
        )
        worksheet.cell(
            row=row_cursor,
            column=headers.get(column_headers.get("annual_sell_price", "Annual Sell Price"), 7),
            value=item.annual_sell_price,
        )
        row_cursor += 1
    workbook.save(target_path)


def _line_items_from_canonical_model(
    *,
    workspace: dict[str, object],
    canonical_model: ClientPricingModelCanonicalArtifact,
) -> list[PricingLineItemArtifact]:
    sites = _site_count(workspace)
    text = _workspace_text(workspace)
    site_multiplier = float(canonical_model.staffing_rules.get("site_multiplier", 1.0) or 1.0)
    continuous_multiplier = float(canonical_model.staffing_rules.get("continuous_coverage_multiplier", 1.35) or 1.35)
    continuous_coverage = "24/7" in text or "seven days" in text or "7 days" in text
    line_items: list[PricingLineItemArtifact] = []
    for category in canonical_model.labor_categories:
        hours_per_week = float(category.default_hours_per_week or 0)
        if category.apply_site_multiplier:
            hours_per_week *= sites * site_multiplier
        if continuous_coverage and category.apply_continuous_coverage_multiplier:
            hours_per_week *= continuous_multiplier
        loaded_hourly_cost = round(category.hourly_rate * category.burden_factor * category.markup_factor, 2)
        annual_sell_price = round(loaded_hourly_cost * hours_per_week * 52, 2)
        line_items.append(
            PricingLineItemArtifact(
                labor_category=category.label,
                hours_per_week=round(hours_per_week, 2),
                hourly_rate=category.hourly_rate,
                burden_factor=category.burden_factor,
                markup_factor=category.markup_factor,
                loaded_hourly_cost=loaded_hourly_cost,
                annual_sell_price=annual_sell_price,
            )
        )
    return line_items


def _write_pricing_workbook(*, target_path: Path, workspace: dict[str, object], line_items: list[PricingLineItemArtifact], assumptions: list[str]) -> None:
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Pricing Summary"
    summary_sheet.append(["Opportunity", workspace.get("opportunity_name", "")])
    summary_sheet.append(["Client", workspace.get("client_name", "")])
    summary_sheet.append([])
    summary_sheet.append([
        "Labor Category",
        "Hours/Week",
        "Hourly Rate",
        "Burden",
        "Markup",
        "Loaded Hourly Cost",
        "Annual Sell Price",
    ])
    for item in line_items:
        summary_sheet.append(
            [
                item.labor_category,
                item.hours_per_week,
                item.hourly_rate,
                item.burden_factor,
                item.markup_factor,
                item.loaded_hourly_cost,
                item.annual_sell_price,
            ]
        )

    assumptions_sheet = workbook.create_sheet("Assumptions")
    assumptions_sheet.append(["Pricing Assumptions"])
    for assumption in assumptions:
        assumptions_sheet.append([assumption])

    workbook.save(target_path)


def build_pricing_package(
    *,
    workspace: dict[str, object],
    target_dir: Path,
    pricing_model: ClientPricingModelResponse | dict[str, Any] | None = None,
    approved_assets: list[dict[str, Any]] | None = None,
) -> PricingPackageArtifact:
    target_dir.mkdir(parents=True, exist_ok=True)
    selected_model = None
    if isinstance(pricing_model, ClientPricingModelResponse):
        selected_model = pricing_model
    elif isinstance(pricing_model, dict) and pricing_model:
        selected_model = ClientPricingModelResponse.model_validate(pricing_model)

    assumptions: list[str] = []
    validation_errors: list[str] = []
    unresolved_pricing_blockers: list[str] = []
    workbook_template_path = None
    line_items: list[PricingLineItemArtifact] = []

    if selected_model and selected_model.canonical_model:
        canonical_model = deepcopy(selected_model.canonical_model)
        line_items = _line_items_from_canonical_model(workspace=workspace, canonical_model=canonical_model)
        assumptions = list(canonical_model.assumptions)
        assumptions.append(f"Pricing uses the approved client model '{selected_model.name}'.")
        assumptions.append("Final margin approval remains a human gate before export.")
        workbook_template_path = selected_model.workbook_template_path
    else:
        tags = _service_tags(workspace)
        pricing_profiles = select_assets(asset_type="pricing_profile", tags=tags, assets=approved_assets)
        if not pricing_profiles:
            return PricingPackageArtifact(
                pricing_narrative="Pricing is blocked because no approved pricing profile is available for this opportunity.",
                unresolved_pricing_blockers=["No approved pricing profile or rate card is available for the current service mix."],
                validation_errors=["Pricing profile missing."],
                approval_required=True,
            )

        pricing_profile = pricing_profiles[0]
        rate_card = pricing_profile.metadata.get("rate_card", {})
        assumptions = [
            f"Pricing uses the approved profile '{pricing_profile.title}'.",
            "Labor hours are derived from deterministic staffing heuristics and should be reviewed by operations.",
            "Final margin approval remains a human gate before export.",
        ]
        for labor_category, hours_per_week in _staffing_template(workspace):
            card = rate_card.get(labor_category)
            if not isinstance(card, dict):
                validation_errors.append(f"Missing rate-card entry for {labor_category}.")
                continue
            hourly_rate = float(card.get("hourly_rate", 0))
            burden_factor = float(card.get("burden_factor", 1))
            markup_factor = float(card.get("markup_factor", 1))
            loaded_hourly_cost = round(hourly_rate * burden_factor * markup_factor, 2)
            annual_sell_price = round(loaded_hourly_cost * hours_per_week * 52, 2)
            line_items.append(
                PricingLineItemArtifact(
                    labor_category=labor_category.replace("_", " ").title(),
                    hours_per_week=hours_per_week,
                    hourly_rate=hourly_rate,
                    burden_factor=burden_factor,
                    markup_factor=markup_factor,
                    loaded_hourly_cost=loaded_hourly_cost,
                    annual_sell_price=annual_sell_price,
                )
            )

    annual_total = round(sum(item.annual_sell_price for item in line_items), 2)
    monthly_total = round(annual_total / 12, 2) if annual_total else 0
    workbook_path = target_dir / "pricing-workbook.xlsx"
    if workbook_template_path:
        try:
            mapping = selected_model.canonical_model.workbook_mapping if selected_model and selected_model.canonical_model else {}
            _write_template_workbook(
                template_path=Path(workbook_template_path),
                target_path=workbook_path,
                line_items=line_items,
                mapping=mapping,
            )
        except Exception:
            validation_errors.append("The selected pricing workbook template could not be populated; the standard workbook was generated instead.")
            _write_pricing_workbook(
                target_path=workbook_path,
                workspace=workspace,
                line_items=line_items,
                assumptions=assumptions,
            )
    else:
        _write_pricing_workbook(
            target_path=workbook_path,
            workspace=workspace,
            line_items=line_items,
            assumptions=assumptions,
        )
    narrative = (
        f"The pricing package uses approved labor categories and controlled rate-card assumptions to support "
        f"an estimated annual sell price of ${annual_total:,.2f} (${monthly_total:,.2f} per month). "
        "This workbook should be reviewed with operations and finance before final export."
    )
    if validation_errors:
        unresolved_pricing_blockers.append("Resolve missing labor-category mappings or template issues in the approved pricing model.")
    return PricingPackageArtifact(
        workbook_path=workbook_path.as_posix(),
        pricing_model_id=selected_model.id if selected_model else None,
        pricing_model_name=selected_model.name if selected_model else None,
        workbook_template_path=workbook_template_path,
        pricing_narrative=narrative,
        staffing_assumptions=assumptions,
        line_items=line_items,
        annual_total=annual_total,
        monthly_total=monthly_total,
        assumptions_log=assumptions,
        validation_errors=validation_errors,
        unresolved_pricing_blockers=unresolved_pricing_blockers,
        approval_required=True,
    )
