import mimetypes
from dataclasses import dataclass
from io import BytesIO


SUPPORTED_UPLOAD_EXTENSIONS = (".txt", ".md", ".pdf", ".docx", ".xlsx", ".pptx")


@dataclass
class BatchDocumentExtractionResult:
    combined_text: str
    parsed_files: list[str]
    skipped_files: list[str]
    warnings: list[str]
    source_filename: str
    documents: list["SourceDocumentExtractionRecord"]


@dataclass
class SourceDocumentExtractionRecord:
    source_filename: str
    content_type: str
    parse_status: str
    skip_reason: str | None
    upload_order: int
    source_size_bytes: int
    extracted_text_length: int
    content_text: str | None
    source_sha256: str | None
    storage_path: str | None
    source_payload: bytes | None


def _guess_content_type(filename: str) -> str:
    guessed, _encoding = mimetypes.guess_type(filename)
    return guessed or "application/octet-stream"


def _normalize_text_value(value: object) -> str:
    text = str(value).replace("\r\n", "\n").replace("\r", "\n").strip()
    lines = [" ".join(line.split()) for line in text.split("\n")]
    return "\n".join(line for line in lines if line)


def _extract_text_from_xlsx_payload(payload: bytes) -> str:
    try:
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("XLSX parsing requires openpyxl. Install openpyxl to parse XLSX files.") from exc

    try:
        workbook = load_workbook(filename=BytesIO(payload), data_only=False, read_only=True)
    except Exception as exc:  # pragma: no cover - exercised via failure behavior, exact library exception varies
        raise RuntimeError("XLSX parsing failed. Ensure the workbook is a valid XLSX file.") from exc

    sheet_sections: list[str] = []
    for sheet in workbook.worksheets:
        row_lines: list[str] = []
        for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            populated_cells: list[str] = []
            for column_index, value in enumerate(row, start=1):
                if value is None:
                    continue
                normalized = _normalize_text_value(value)
                if not normalized:
                    continue
                populated_cells.append(f"{get_column_letter(column_index)}{row_index}={normalized}")
            if populated_cells:
                row_lines.append(" | ".join(populated_cells))
        if row_lines:
            sheet_sections.append(f"Sheet: {sheet.title}\n" + "\n".join(row_lines))
        else:
            sheet_sections.append(f"Sheet: {sheet.title}\n[blank sheet]")

    workbook.close()
    return "\n\n".join(sheet_sections)


def _extract_pptx_notes(slide) -> str:
    if not getattr(slide, "has_notes_slide", False):
        return ""

    notes: list[str] = []
    for shape in slide.notes_slide.shapes:
        text = _normalize_text_value(getattr(shape, "text", ""))
        if not text:
            continue
        if text.lower() == "click to add notes":
            continue
        if text in notes:
            continue
        notes.append(text)
    return "\n".join(notes)


def _extract_text_from_pptx_payload(payload: bytes) -> str:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise RuntimeError("PPTX parsing requires python-pptx. Install python-pptx to parse PPTX files.") from exc

    try:
        presentation = Presentation(BytesIO(payload))
    except Exception as exc:  # pragma: no cover - exercised via failure behavior, exact library exception varies
        raise RuntimeError("PPTX parsing failed. Ensure the slide deck is a valid PPTX file.") from exc

    slide_sections: list[str] = []
    for slide_index, slide in enumerate(presentation.slides, start=1):
        content_lines: list[str] = []
        for shape in slide.shapes:
            text = _normalize_text_value(getattr(shape, "text", ""))
            if not text:
                continue
            if text in content_lines:
                continue
            content_lines.append(text)

        notes = _extract_pptx_notes(slide)
        section_lines = [f"Slide {slide_index}"]
        if content_lines:
            section_lines.extend(content_lines)
        if notes:
            section_lines.append("Speaker notes:")
            section_lines.append(notes)
        slide_sections.append("\n".join(section_lines))

    return "\n\n".join(slide_sections)


def _build_batch_result(documents: list["SourceDocumentExtractionRecord"]) -> BatchDocumentExtractionResult:
    parsed_files = [item.source_filename for item in documents if item.parse_status == "PARSED"]
    skipped_files = [item.source_filename for item in documents if item.parse_status == "SKIPPED"]
    warnings = [
        f"{item.source_filename}: {item.skip_reason}"
        for item in documents
        if item.parse_status == "SKIPPED" and item.skip_reason
    ]
    combined_parts = [
        f"Source file: {item.source_filename}\n{item.content_text.strip()}"
        for item in documents
        if item.parse_status == "PARSED" and item.content_text and item.content_text.strip()
    ]
    parsed_count = len(parsed_files)
    return BatchDocumentExtractionResult(
        combined_text="\n\n".join(combined_parts),
        parsed_files=parsed_files,
        skipped_files=skipped_files,
        warnings=warnings,
        source_filename=f"intake-batch-{parsed_count}-files.txt",
        documents=documents,
    )


def extract_text_from_manual_text(
    filename: str,
    raw_text: str,
    *,
    upload_order: int = 1,
) -> SourceDocumentExtractionRecord:
    normalized = raw_text.strip()
    payload = raw_text.encode("utf-8")
    if len(normalized) < 20:
        return SourceDocumentExtractionRecord(
            source_filename=filename,
            content_type="text/plain",
            parse_status="SKIPPED",
            skip_reason="extracted content is too short to parse.",
            upload_order=upload_order,
            source_size_bytes=len(payload),
            extracted_text_length=len(normalized),
            content_text=None,
            source_sha256=None,
            storage_path=None,
            source_payload=payload,
        )

    return SourceDocumentExtractionRecord(
        source_filename=filename,
        content_type="text/plain",
        parse_status="PARSED",
        skip_reason=None,
        upload_order=upload_order,
        source_size_bytes=len(payload),
        extracted_text_length=len(normalized),
        content_text=normalized,
        source_sha256=None,
        storage_path=None,
        source_payload=payload,
    )


def extract_text_from_mixed_inputs(
    files: list[tuple[str, bytes]] | None = None,
    *,
    text_entries: list[tuple[str, str]] | None = None,
) -> BatchDocumentExtractionResult:
    documents: list[SourceDocumentExtractionRecord] = []
    upload_order = 1

    for filename, payload in files or []:
        content_type = _guess_content_type(filename)
        try:
            content = extract_text_from_upload(filename, payload)
        except RuntimeError as exc:
            documents.append(
                SourceDocumentExtractionRecord(
                    source_filename=filename,
                    content_type=content_type,
                    parse_status="SKIPPED",
                    skip_reason=str(exc),
                    upload_order=upload_order,
                    source_size_bytes=len(payload),
                    extracted_text_length=0,
                    content_text=None,
                    source_sha256=None,
                    storage_path=None,
                    source_payload=payload,
                )
            )
            upload_order += 1
            continue

        normalized = content.strip()
        if len(normalized) < 20:
            documents.append(
                SourceDocumentExtractionRecord(
                    source_filename=filename,
                    content_type=content_type,
                    parse_status="SKIPPED",
                    skip_reason="extracted content is too short to parse.",
                    upload_order=upload_order,
                    source_size_bytes=len(payload),
                    extracted_text_length=len(normalized),
                    content_text=None,
                    source_sha256=None,
                    storage_path=None,
                    source_payload=payload,
                )
            )
            upload_order += 1
            continue

        documents.append(
            SourceDocumentExtractionRecord(
                source_filename=filename,
                content_type=content_type,
                parse_status="PARSED",
                skip_reason=None,
                upload_order=upload_order,
                source_size_bytes=len(payload),
                extracted_text_length=len(normalized),
                content_text=normalized,
                source_sha256=None,
                storage_path=None,
                source_payload=payload,
            )
        )
        upload_order += 1

    for filename, raw_text in text_entries or []:
        documents.append(extract_text_from_manual_text(filename, raw_text, upload_order=upload_order))
        upload_order += 1

    return _build_batch_result(documents)


def extract_text_from_upload(filename: str, payload: bytes) -> str:
    lowered = filename.lower()
    if lowered.endswith(".txt") or lowered.endswith(".md"):
        return payload.decode("utf-8", errors="ignore")

    if lowered.endswith(".pdf"):
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("PDF parsing requires pypdf. Install pypdf to parse PDF files.") from exc

        reader = PdfReader(BytesIO(payload))
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    if lowered.endswith(".docx"):
        try:
            import docx
        except ImportError as exc:
            raise RuntimeError("DOCX parsing requires python-docx. Install python-docx.") from exc

        document = docx.Document(BytesIO(payload))
        return "\n".join(paragraph.text for paragraph in document.paragraphs)

    if lowered.endswith(".xlsx"):
        return _extract_text_from_xlsx_payload(payload)

    if lowered.endswith(".pptx"):
        return _extract_text_from_pptx_payload(payload)

    raise RuntimeError("Unsupported file type. Use TXT, MD, PDF, DOCX, XLSX, or PPTX.")


def extract_text_from_upload_batch(files: list[tuple[str, bytes]]) -> BatchDocumentExtractionResult:
    return extract_text_from_mixed_inputs(files)
