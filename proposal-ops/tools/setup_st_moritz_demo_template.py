from __future__ import annotations

import json
import shutil
import subprocess
import sys
import textwrap
import time
from pathlib import Path
from urllib.request import urlopen

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image
from sqlalchemy import select


PROPOSAL_OPS_DIR = Path(__file__).resolve().parents[1]
WORKSPACE_ROOT = PROPOSAL_OPS_DIR.parent
if str(PROPOSAL_OPS_DIR) not in sys.path:
    sys.path.insert(0, str(PROPOSAL_OPS_DIR))

from app.core.db import SessionLocal  # noqa: E402
from app.modules.proposal_builder.client_profiles import ClientProfileService  # noqa: E402
from app.modules.proposal_builder.export import build_export_bundle  # noqa: E402
from app.modules.proposal_builder.models import ClientProfile  # noqa: E402
from app.modules.proposal_builder.schemas import ClientProfileCreateRequest, ClientProfileDefaultsRequest  # noqa: E402
from app.modules.proposal_builder.template_renderer import _append_page_number_run, validate_proposal_template  # noqa: E402


LOGO_URL = "https://smssi.com/wp-content/themes/Divi-weCreate-Child/images/smssi-logo.svg"
HOME_URL = "https://smssi.com/"
ABOUT_URL = "https://smssi.com/about/"
SERVICES_URL = "https://smssi.com/services/"
TECHNOLOGY_URL = "https://smssi.com/technology/"

BRAND_RED = "D0001C"
BRAND_DEEP_RED = "A20502"
BRAND_DARK_RED = "5A1111"
BRAND_NAVY = "1A2536"
BRAND_GOLD = "B57420"
BRAND_LIGHT_GRAY = "E7E9EC"

CLIENT_PROFILE_NAME = "St. Moritz Security Services, Inc."
TEMPLATE_NAME = "St. Moritz Co-Branded Proposal Shell"
TEMPLATE_DIR = PROPOSAL_OPS_DIR / "demo-artifacts" / "client-templates" / "st-moritz"
SVG_PATH = TEMPLATE_DIR / "st-moritz-logo.svg"
RAW_LOGO_SCREENSHOT_PATH = TEMPLATE_DIR / "st-moritz-logo-raw.png"
PNG_PATH = TEMPLATE_DIR / "st-moritz-logo.png"
TEMPLATE_PATH = TEMPLATE_DIR / "st-moritz-co-branded-rfp-template.docx"
SAMPLE_EXPORT_DIR = TEMPLATE_DIR / "sample-export"
SAMPLE_MANIFEST_PATH = TEMPLATE_DIR / "st-moritz-demo-template-manifest.json"
SAMPLE_NOTES_PATH = TEMPLATE_DIR / "st-moritz-demo-template-notes.md"

SECTION_SPECS = [
    ("Transmittal Letter", "transmittal_letter"),
    ("Executive Summary", "executive_summary"),
    ("Technical Approach", "technical_approach"),
    ("Management and Staffing Plan", "management_staffing_plan"),
    ("Transition and Mobilization Plan", "transition_mobilization_plan"),
    ("Past Performance", "past_performance"),
    ("Pricing and Commercials", "pricing_and_commercials"),
    ("Compliance and Attachments", "compliance_attachments"),
]


def _find_edge() -> Path:
    candidates = [
        shutil.which("msedge"),
        shutil.which("msedge.exe"),
        r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
    ]
    for candidate in candidates:
        if candidate and Path(candidate).exists():
            return Path(candidate)
    raise RuntimeError("Microsoft Edge was not found. The logo PNG cannot be generated.")


def _download_logo_svg() -> None:
    TEMPLATE_DIR.mkdir(parents=True, exist_ok=True)
    SVG_PATH.write_bytes(urlopen(LOGO_URL).read())


def _write_logo_capture_html(path: Path) -> None:
    path.write_text(
        textwrap.dedent(
            f"""
            <!doctype html>
            <html lang="en">
            <head>
              <meta charset="utf-8">
              <style>
                html, body {{
                  margin: 0;
                  padding: 0;
                  background: white;
                  width: 1100px;
                  height: 260px;
                  overflow: hidden;
                }}
                body {{
                  display: grid;
                  place-items: center;
                }}
                img {{
                  width: 760px;
                  height: auto;
                  display: block;
                }}
              </style>
            </head>
            <body>
              <img src="{SVG_PATH.as_uri()}" alt="St. Moritz Security Services, Inc.">
            </body>
            </html>
            """
        ).strip(),
        encoding="utf-8",
    )


def _make_white_transparent(image: Image.Image) -> Image.Image:
    rgba = image.convert("RGBA")
    pixels = rgba.load()
    for y in range(rgba.height):
        for x in range(rgba.width):
            red, green, blue, alpha = pixels[x, y]
            if red >= 246 and green >= 246 and blue >= 246:
                pixels[x, y] = (255, 255, 255, 0)
    return rgba


def _crop_to_content(image: Image.Image, padding: int = 12) -> Image.Image:
    rgba = image.convert("RGBA")
    alpha = rgba.getchannel("A")
    bbox = alpha.getbbox()
    if not bbox:
        return rgba
    left = max(bbox[0] - padding, 0)
    top = max(bbox[1] - padding, 0)
    right = min(bbox[2] + padding, rgba.width)
    bottom = min(bbox[3] + padding, rgba.height)
    return rgba.crop((left, top, right, bottom))


def _render_logo_png() -> None:
    edge = _find_edge()
    capture_html = TEMPLATE_DIR / "st-moritz-logo-capture.html"
    _write_logo_capture_html(capture_html)
    subprocess.run(
        [
            str(edge),
            "--headless=new",
            "--disable-gpu",
            f"--screenshot={RAW_LOGO_SCREENSHOT_PATH}",
            "--window-size=1100,260",
            capture_html.as_uri(),
        ],
        check=True,
        capture_output=True,
        text=True,
    )
    for _ in range(40):
        if RAW_LOGO_SCREENSHOT_PATH.exists():
            break
        time.sleep(0.25)
    if not RAW_LOGO_SCREENSHOT_PATH.exists():
        raise RuntimeError("Edge did not write the St. Moritz logo screenshot.")
    raw_image = Image.open(RAW_LOGO_SCREENSHOT_PATH)
    transparent = _make_white_transparent(raw_image)
    cropped = _crop_to_content(transparent)
    cropped.save(PNG_PATH)


def _set_cell_fill(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _set_cell_border(cell, *, bottom: str | None = None) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    if bottom:
        edge = OxmlElement("w:bottom")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "12")
        edge.set(qn("w:space"), "0")
        edge.set(qn("w:color"), bottom)
        borders.append(edge)


def _set_page_margins(document: Document) -> None:
    for section in document.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.65)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        section.different_first_page_header_footer = True


def _set_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    title = document.styles["Title"]
    title.font.name = "Arial Narrow"
    title.font.bold = True
    title.font.size = Pt(26)
    title.font.color.rgb = RGBColor.from_string(BRAND_NAVY)

    heading1 = document.styles["Heading 1"]
    heading1.font.name = "Arial Narrow"
    heading1.font.bold = True
    heading1.font.size = Pt(17)
    heading1.font.color.rgb = RGBColor.from_string(BRAND_NAVY)

    heading2 = document.styles["Heading 2"]
    heading2.font.name = "Arial Narrow"
    heading2.font.bold = True
    heading2.font.size = Pt(12)
    heading2.font.color.rgb = RGBColor.from_string(BRAND_DEEP_RED)


def _add_top_band(document: Document) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    _set_cell_fill(cell, BRAND_RED)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("ST. MORITZ | PROPOSAL RESPONSE SHELL")
    run.font.name = "Arial Narrow"
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor(255, 255, 255)


def _add_cover_page(document: Document) -> None:
    _add_top_band(document)
    document.add_paragraph()

    logo_paragraph = document.add_paragraph()
    logo_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo_run = logo_paragraph.add_run()
    logo_run.add_picture(str(PNG_PATH), width=Inches(3.65))

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Enterprise Proposal Response Shell")

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Co-branded response format for security, technology, and business-services opportunities")
    run.font.name = "Calibri"
    run.font.size = Pt(11.5)
    run.font.color.rgb = RGBColor.from_string(BRAND_DARK_RED)

    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = note.add_run("Prepared by Boss Key | Solution Architecture and Proposal Automation Support")
    run.font.name = "Calibri"
    run.font.size = Pt(10.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(BRAND_NAVY)

    document.add_paragraph()

    metadata = document.add_table(rows=4, cols=2)
    metadata.alignment = WD_TABLE_ALIGNMENT.CENTER
    rows = [
        ("Prepared For", "{{ client_name }}"),
        ("Opportunity", "{{ opportunity_name }}"),
        ("Solicitation / Reference", "{{ solicitation_number }}"),
        ("Proposal Due", "{{ proposal_due }}"),
    ]
    for index, (label, value) in enumerate(rows):
        left = metadata.rows[index].cells[0]
        right = metadata.rows[index].cells[1]
        left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _set_cell_fill(left, BRAND_NAVY)
        _set_cell_fill(right, BRAND_LIGHT_GRAY)
        left_run = left.paragraphs[0].add_run(label)
        left_run.font.name = "Arial Narrow"
        left_run.font.bold = True
        left_run.font.size = Pt(11)
        left_run.font.color.rgb = RGBColor(255, 255, 255)
        right_run = right.paragraphs[0].add_run(value)
        right_run.font.name = "Calibri"
        right_run.font.size = Pt(11)
        right_run.font.color.rgb = RGBColor.from_string(BRAND_NAVY)
        _set_cell_border(right, bottom=BRAND_GOLD)

    document.add_paragraph()
    lead_in = document.add_paragraph()
    lead_in.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lead_run = lead_in.add_run(
        "Designed to show St. Moritz branding, section structure, metadata handling, and export-ready formatting in the Proposal Builder."
    )
    lead_run.font.name = "Calibri"
    lead_run.font.size = Pt(10)
    lead_run.italic = True
    lead_run.font.color.rgb = RGBColor.from_string(BRAND_NAVY)
    document.add_page_break()


def _add_footer(document: Document) -> None:
    section = document.sections[0]
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    left = paragraph.add_run("{{ client_name }} | {{ solicitation_number }} | Due {{ proposal_due }} | ")
    left.font.name = "Calibri"
    left.font.size = Pt(8.5)
    left.font.color.rgb = RGBColor.from_string(BRAND_NAVY)
    _append_page_number_run(paragraph, prefix="Page ")
    for run in paragraph.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor.from_string(BRAND_NAVY)


def _add_section_shell(document: Document, title: str, anchor: str, *, first: bool = False) -> None:
    if not first:
        document.add_page_break()

    heading_table = document.add_table(rows=1, cols=1)
    heading_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = heading_table.rows[0].cells[0]
    _set_cell_fill(cell, BRAND_RED)
    _set_cell_border(cell, bottom=BRAND_GOLD)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(title.upper())
    run.font.name = "Arial Narrow"
    run.font.size = Pt(15)
    run.font.bold = True
    run.font.color.rgb = RGBColor(255, 255, 255)

    descriptor = document.add_paragraph()
    lead = descriptor.add_run("St. Moritz visual identity leads. ")
    lead.font.name = "Calibri"
    lead.font.size = Pt(10)
    lead.font.bold = True
    lead.font.color.rgb = RGBColor.from_string(BRAND_NAVY)
    tail = descriptor.add_run(
        "This section is populated at export time from the Proposal Builder and stays reusable across future proposal responses."
    )
    tail.font.name = "Calibri"
    tail.font.size = Pt(10)
    tail.font.color.rgb = RGBColor.from_string(BRAND_NAVY)

    accent = document.add_paragraph()
    accent_run = accent.add_run("Prepared by Boss Key for demonstration and client deployment planning.")
    accent_run.font.name = "Calibri"
    accent_run.font.size = Pt(9.5)
    accent_run.italic = True
    accent_run.font.color.rgb = RGBColor.from_string(BRAND_DARK_RED)

    document.add_paragraph()
    document.add_paragraph(f"{{{{ {anchor} }}}}")


def _build_template() -> None:
    document = Document()
    _set_page_margins(document)
    _set_styles(document)
    document.core_properties.title = "St. Moritz Co-Branded Proposal Shell"
    document.core_properties.subject = "St. Moritz x Boss Key proposal response template"
    document.core_properties.author = "Boss Key LLC"
    document.core_properties.company = "Boss Key LLC"
    document.core_properties.comments = "Built for St. Moritz Security Services, Inc. demo use."

    _add_cover_page(document)
    _add_footer(document)

    for index, (title, anchor) in enumerate(SECTION_SPECS):
        _add_section_shell(document, title, anchor, first=index == 0)

    TEMPLATE_DIR.mkdir(parents=True, exist_ok=True)
    document.save(TEMPLATE_PATH)


def _demo_workspace() -> dict[str, object]:
    return {
        "opportunity_id": "st-moritz-demo-opp",
        "opportunity_name": "Proposal Operations and RFP Response Enablement",
        "client_name": CLIENT_PROFILE_NAME,
        "solicitation_id": "SMSSI-DEMO-2026-PR-01",
        "opportunity_summary": {
            "client_name": CLIENT_PROFILE_NAME,
            "opportunity_name": "Proposal Operations and RFP Response Enablement",
            "solicitation_number": "SMSSI-DEMO-2026-PR-01",
            "proposal_due_date": "March 31, 2026",
            "proposal_due_time": "5:00 PM ET",
            "strategic_fit": "Strong fit for St. Moritz growth, capture, and proposal modernization priorities.",
            "geographic_fit": "Designed to support St. Moritz regional and national pursuits.",
            "operational_fit": "Matches cross-functional sales, operations, legal, finance, and technology review needs.",
            "revenue_scale_signal": "Commercial structure can be phased from pilot to managed support.",
            "complexity_level": "Moderate",
            "risk_level": "Medium",
            "recommended_next_action": "Use the shell to align on output quality, section logic, and deployment expectations.",
        },
        "client_sections": [
            {
                "section_title": "Transmittal Letter",
                "body_markdown": textwrap.dedent(
                    """
                    [Date]

                    Eric Viletto
                    Vice President of Sales
                    St. Moritz Security Services, Inc.

                    Boss Key is pleased to provide this co-branded proposal response shell for St. Moritz Security Services, Inc. The purpose of this document is to demonstrate how a live RFP can be translated into a branded, review-ready response package without forcing the proposal team into manual restructuring work each time.

                    This shell is intentionally reusable across St. Moritz opportunities in security services, remote services, technology-enabled offerings, and related business-service pursuits.
                    """
                ).strip(),
                "used_content_block_ids": [],
                "cited_requirement_codes": [],
            },
            {
                "section_title": "Executive Summary",
                "body_markdown": textwrap.dedent(
                    """
                    St. Moritz presents as a high-touch operator with a national footprint, combining local office support with centralized capabilities, disciplined training, and responsive service management. This response shell is designed to carry that identity consistently into competitive proposals.

                    The structure supports three things that matter immediately in RFP response work:
                    - a clear executive narrative aligned to buyer priorities
                    - section-level control over technical, staffing, transition, pricing, and compliance content
                    - export-ready formatting that reduces hand rework before final production in Word

                    In the demo environment, Boss Key uses this shell to show how St. Moritz can move from raw solicitation material to a branded proposal package starter with visible rigor, section discipline, and a cleaner path to final submission.
                    """
                ).strip(),
                "used_content_block_ids": [],
                "cited_requirement_codes": [],
            },
            {
                "section_title": "Technical Approach",
                "body_markdown": textwrap.dedent(
                    """
                    The proposed operating model starts with reliable intake of the RFP package and continues through structured extraction, compliance mapping, narrative planning, and export packaging. For St. Moritz, that means the response process can be built around real operational themes instead of disconnected manual documents.

                    ### Core response workflow
                    - ingest the base solicitation and amendments into one opportunity record
                    - extract dates, submission instructions, evaluation criteria, and mandatory requirements
                    - map those requirements into a compliance matrix and response outline
                    - generate first-pass narrative sections tied to approved positioning and section logic
                    - export a Word-ready package the team can complete and finalize

                    The result is a response workflow that supports consistency, faster first drafts, and clearer collaboration across capture, operations, and leadership stakeholders.
                    """
                ).strip(),
                "used_content_block_ids": [],
                "cited_requirement_codes": [],
            },
            {
                "section_title": "Management and Staffing Plan",
                "body_markdown": textwrap.dedent(
                    """
                    St. Moritz already emphasizes local office relationships, responsive management, and centralized operational support through its broader platform. This shell is meant to preserve that management story in proposal form while giving the team a more structured content-production process.

                    The proposal operating model can support role-specific review and contribution across:
                    - sales leadership and capture ownership
                    - operations and service-delivery subject matter experts
                    - legal and compliance reviewers
                    - pricing and finance stakeholders
                    - technology and deployment reviewers when a solution includes remote services, SOC, or systems integration

                    That structure gives St. Moritz a clearer path from opportunity intake to controlled final draft without losing accountability for who owns each section.
                    """
                ).strip(),
                "used_content_block_ids": [],
                "cited_requirement_codes": [],
            },
            {
                "section_title": "Transition and Mobilization Plan",
                "body_markdown": textwrap.dedent(
                    """
                    Boss Key's recommended first step remains a controlled pilot so the proposal workflow can be validated against live St. Moritz materials. The two-week pilot is intended to prove output quality, collaboration flow, and deployment fit before broader rollout.

                    ### Pilot mobilization focus
                    - confirm security, hosting, and access constraints
                    - run one or two representative RFP packages through the intake and digest flow
                    - review outputs with stakeholders and tune structure, section logic, and review expectations
                    - finalize the production-ready operating approach for ongoing use

                    This staged rollout keeps the implementation practical while showing St. Moritz how the system can mature into a repeatable internal proposal capability.
                    """
                ).strip(),
                "used_content_block_ids": [],
                "cited_requirement_codes": [],
            },
            {
                "section_title": "Past Performance",
                "body_markdown": textwrap.dedent(
                    """
                    In a live deployment, this section would draw from approved St. Moritz reference material, comparable engagements, and sector-specific proof points. For demo purposes, the important point is that the shell has a dedicated location for evaluator-facing experience and credibility evidence rather than forcing the team to improvise format each time.

                    Recommended future insertions include:
                    - approved transportation, corporate, healthcare, or logistics security references
                    - relevant remote services or technology-enabled operations examples
                    - quality-control, training, and management proof points
                    - sector-aligned differentiators that support buyer confidence
                    """
                ).strip(),
                "used_content_block_ids": [],
                "cited_requirement_codes": [],
            },
            {
                "section_title": "Pricing and Commercials",
                "body_markdown": textwrap.dedent(
                    """
                    The commercial model already discussed with St. Moritz can be expressed cleanly in this shell while keeping pricing logic separate from the final technical narrative where required.

                    ### Current demo commercial structure
                    - 2-week pilot aligned to a fixed diagnostic engagement
                    - optional managed monthly support for higher proposal volume
                    - optional lighter advisory support when the operating need is narrower

                    The template keeps the pricing section organized and presentation-ready, while allowing client-specific workbook logic and commercial terms to be inserted during final packaging.
                    """
                ).strip(),
                "used_content_block_ids": [],
                "cited_requirement_codes": [],
            },
            {
                "section_title": "Compliance and Attachments",
                "body_markdown": textwrap.dedent(
                    """
                    This final section is reserved for the items that often make or break submission readiness: forms, acknowledgements, attachments, deployment assumptions, and compliance notes.

                    For St. Moritz, the demo shell is designed to accommodate:
                    - NDA and scoping documentation
                    - deployment and security review notes
                    - client-required forms or workbooks
                    - final attachments, resumes, references, and certifications
                    - proposal-specific compliance tables and submission controls

                    Keeping these controls in a dedicated section helps the final Word package feel orderly, auditable, and easier to review internally before submission.
                    """
                ).strip(),
                "used_content_block_ids": [],
                "cited_requirement_codes": [],
            },
        ],
        "draft_package": {
            "sections": [],
            "client_sections": [],
            "internal_notes": [],
            "unresolved_items": [],
            "editor_notes": "Demo collateral generated from the St. Moritz co-branded template.",
        },
        "compliance_matrix": [
            {
                "requirement_code": "SM-001",
                "requirement_text": "Use a consistent branded shell for executive, technical, staffing, transition, pricing, and compliance content.",
                "requirement_type": "DEMO_REQUIRED",
                "proposal_section": "Executive Summary",
                "owner": "Proposal Lead",
                "status": "READY",
            },
            {
                "requirement_code": "SM-002",
                "requirement_text": "Preserve St. Moritz positioning around responsive service, national reach, and technology-enabled operations.",
                "requirement_type": "DEMO_REQUIRED",
                "proposal_section": "Technical Approach",
                "owner": "Solution Architect",
                "status": "READY",
            },
        ],
        "win_themes": [],
        "proposal_outline": {"sections": []},
    }


def _sample_notes() -> str:
    return textwrap.dedent(
        f"""
        # St. Moritz Demo Proposal Template

        This asset set is modeled after current public St. Moritz branding and positioning:

        - Home: {HOME_URL}
        - About: {ABOUT_URL}
        - Services: {SERVICES_URL}
        - Technology: {TECHNOLOGY_URL}
        - Public logo source: {LOGO_URL}

        ## Brand cues applied

        - Primary red: #{BRAND_RED}
        - Secondary red: #{BRAND_DEEP_RED}
        - Dark red: #{BRAND_DARK_RED}
        - Navy: #{BRAND_NAVY}
        - Gold accent: #{BRAND_GOLD}
        - Light gray: #{BRAND_LIGHT_GRAY}
        - Heading font fallback: Arial Narrow
        - Body font fallback: Calibri

        ## Delivered files

        - Co-branded template: {TEMPLATE_PATH}
        - Rendered logo PNG: {PNG_PATH}
        - Sample rendered collateral export: {SAMPLE_EXPORT_DIR}
        """
    ).strip() + "\n"


def _ensure_st_moritz_profile_and_template() -> dict[str, object]:
    with SessionLocal() as db:
        service = ClientProfileService(db)
        profile = db.scalar(select(ClientProfile).where(ClientProfile.display_name == CLIENT_PROFILE_NAME))
        if profile is None:
            created = service.create_profile(
                ClientProfileCreateRequest(
                    display_name=CLIENT_PROFILE_NAME,
                    aliases=["St Moritz", "St. Moritz", "SMSSI"],
                    approved_content_tags=["security", "remote-services", "technology", "proposal-demo"],
                    region_tags=["national"],
                    service_tags=["security-services", "remote-services", "technology"],
                )
            )
            profile_id = created.id
            default_pricing_model_id = created.default_pricing_model_id
        else:
            profile.aliases_json = json.dumps(["St Moritz", "St. Moritz", "SMSSI"])
            profile.approved_content_tags_json = json.dumps(["security", "remote-services", "technology", "proposal-demo"])
            profile.region_tags_json = json.dumps(["national"])
            profile.service_tags_json = json.dumps(["security-services", "remote-services", "technology"])
            profile.status = "ACTIVE"
            db.commit()
            profile_id = profile.id
            default_pricing_model_id = profile.default_pricing_model_id

        template_response = service.upload_proposal_template(
            profile_id,
            name=TEMPLATE_NAME,
            filename=TEMPLATE_PATH.name,
            content=TEMPLATE_PATH.read_bytes(),
        )
        profile_response = service.set_defaults(
            profile_id,
            ClientProfileDefaultsRequest(
                default_pricing_model_id=default_pricing_model_id,
                default_proposal_template_id=template_response.id,
            ),
        )
        return {
            "profile_id": profile_response.id,
            "profile_name": profile_response.display_name,
            "template_id": template_response.id,
            "template_name": template_response.name,
            "template_status": template_response.status,
            "default_proposal_template_id": profile_response.default_proposal_template_id,
        }


def _build_sample_export(template_metadata: dict[str, object]) -> dict[str, object]:
    SAMPLE_EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    manifest = build_export_bundle(
        _demo_workspace(),
        proposal_template={
            "id": template_metadata["template_id"],
            "name": template_metadata["template_name"],
            "storage_path": TEMPLATE_PATH.as_posix(),
        },
    )
    copied_docx = SAMPLE_EXPORT_DIR / "st-moritz-demo-collateral.docx"
    copied_zip = SAMPLE_EXPORT_DIR / "st-moritz-demo-support-bundle.zip"
    shutil.copy2(manifest["client_docx_path"], copied_docx)
    shutil.copy2(manifest["support_bundle_zip_path"], copied_zip)
    manifest["copied_client_docx_path"] = copied_docx.as_posix()
    manifest["copied_support_bundle_zip_path"] = copied_zip.as_posix()
    return manifest


def main() -> None:
    _download_logo_svg()
    _render_logo_png()
    _build_template()
    validation = validate_proposal_template(TEMPLATE_PATH)
    if validation["status"] != "VALID":
        raise RuntimeError(f"Generated template failed validation: {validation}")

    template_metadata = _ensure_st_moritz_profile_and_template()
    export_manifest = _build_sample_export(template_metadata)

    TEMPLATE_DIR.mkdir(parents=True, exist_ok=True)
    SAMPLE_NOTES_PATH.write_text(_sample_notes(), encoding="utf-8")
    SAMPLE_MANIFEST_PATH.write_text(
        json.dumps(
            {
                "template": {
                    "path": TEMPLATE_PATH.as_posix(),
                    "validation": validation,
                    "logo_svg_path": SVG_PATH.as_posix(),
                    "logo_png_path": PNG_PATH.as_posix(),
                },
                "client_profile": template_metadata,
                "sample_export": export_manifest,
            },
            indent=2,
            default=str,
        ),
        encoding="utf-8",
    )

    print(f"Template created: {TEMPLATE_PATH}")
    print(f"Sample export: {export_manifest['client_docx_path']}")
    print(f"Client profile: {template_metadata['profile_name']} ({template_metadata['profile_id']})")
    print(f"Template ID: {template_metadata['template_id']}")
    print(f"Manifest: {SAMPLE_MANIFEST_PATH}")


if __name__ == "__main__":
    main()
